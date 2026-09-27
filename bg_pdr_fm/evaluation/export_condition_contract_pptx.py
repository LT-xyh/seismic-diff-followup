"""Export the compact Figure 3 diagnostics as native editable PowerPoint shapes."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any, Mapping, Sequence

import numpy as np
from scipy.stats import gaussian_kde

from bg_pdr_fm.evaluation.visualize_residual_transport import SUBSET_ORDER


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_DIR = REPO_ROOT / "docs" / "paper" / "AAAI2027" / "figures" / "figure3_condition_learning"

SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 5.72
FONT_NAME = "Times New Roman"

GREEN = "178F48"
GREEN_LIGHT = "DCEFE4"
BLUE = "1E63D5"
BLUE_LIGHT = "DCE8F8"
PURPLE = "8357A6"
PURPLE_LIGHT = "E9E0F0"
GRAY = "8A929B"
GRAY_LIGHT = "D5D9DE"
TEXT = "252A31"
MUTED = "596270"
AXIS = "4C5563"
GRID = "E1E5EA"

ANCHOR_LEFT = 0.82
ANCHOR_TOP = 0.83
ANCHOR_WIDTH = 3.93
ANCHOR_HEIGHT = 3.75
ANCHOR_Y_MIN = -0.55
ANCHOR_Y_MAX = 1.05

ROLE_LEFT = 5.72
ROLE_WIDTH = 6.98
ROLE_HEIGHT = 1.02
ROLE_TOPS = (0.82, 2.30, 3.78)


def _resolve_path(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else REPO_ROOT / path


def _read_csv(path: str | Path) -> list[dict[str, str]]:
    resolved = _resolve_path(path)
    if not resolved.is_file():
        raise FileNotFoundError(f"CSV does not exist: {resolved}")
    with resolved.open(newline="", encoding="utf-8") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _rgb(hex_value: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_value)


def _add_text(
    slide: Any,
    text: str,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    size: float,
    color: str = TEXT,
    bold: bool = False,
    italic: bool = False,
    align: str = "left",
    rotation: float = 0.0,
    name: str | None = None,
) -> Any:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if name:
        shape.name = name
    shape.rotation = rotation
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }[align]
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)
    return shape


def _add_line(
    slide: Any,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str,
    width: float = 0.7,
    dashed: bool = False,
    name: str | None = None,
) -> Any:
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches, Pt

    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    if name:
        line.name = name
    line.line.color.rgb = _rgb(color)
    line.line.width = Pt(width)
    if dashed:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return line


def _add_circle(slide: Any, *, left: float, top: float, diameter: float, color: str, name: str) -> Any:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(left),
        Inches(top),
        Inches(diameter),
        Inches(diameter),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.color.rgb = _rgb(color)
    shape.line.width = Pt(0.5)
    return shape


def _anchor_y(value: float) -> float:
    return ANCHOR_TOP + (ANCHOR_Y_MAX - float(value)) / (ANCHOR_Y_MAX - ANCHOR_Y_MIN) * ANCHOR_HEIGHT


def _role_y(value: float, top: float) -> float:
    return top + (1.0 - float(value)) * ROLE_HEIGHT


def _add_violin(
    slide: Any,
    values: np.ndarray,
    *,
    center_x: float,
    max_half_width: float,
    fill_color: str,
    line_color: str,
    name: str,
) -> Any:
    from pptx.util import Inches, Pt

    values = np.asarray(values, dtype=np.float64)
    low = max(ANCHOR_Y_MIN, float(values.min()) - 0.025)
    high = min(ANCHOR_Y_MAX, float(values.max()) + 0.025)
    grid = np.linspace(low, high, 150)
    if values.size < 2 or float(values.std()) < 1e-10:
        scale = max((high - low) / 6.0, 1e-3)
        density = np.exp(-0.5 * np.square((grid - float(values.mean())) / scale))
    else:
        density = gaussian_kde(values, bw_method="scott")(grid)
    density /= max(float(density.max()), 1e-12)
    widths = max_half_width * density
    points = [
        *((center_x - float(width), _anchor_y(value)) for width, value in zip(widths, grid)),
        *((center_x + float(width), _anchor_y(value)) for width, value in zip(widths[::-1], grid[::-1])),
    ]
    local_scale = float(Inches(1)) / 1000.0
    local_points = [(int(round(x * 1000)), int(round(y * 1000))) for x, y in points]
    builder = slide.shapes.build_freeform(
        start_x=local_points[0][0],
        start_y=local_points[0][1],
        scale=(local_scale, local_scale),
    )
    builder.add_line_segments(local_points[1:], close=True)
    shape = builder.convert_to_shape()
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill_color)
    shape.line.color.rgb = _rgb(line_color)
    shape.line.width = Pt(0.75)
    return shape


def _add_role_boxplot(
    slide: Any,
    values: np.ndarray,
    *,
    center_x: float,
    top: float,
    category_width: float,
    line_color: str,
    fill_color: str,
    name: str,
) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    values = np.asarray(values, dtype=np.float64)
    q1, median, q3 = np.quantile(values, [0.25, 0.50, 0.75])
    iqr = q3 - q1
    lower_candidates = values[values >= q1 - 1.5 * iqr]
    upper_candidates = values[values <= q3 + 1.5 * iqr]
    lower = float(lower_candidates.min())
    upper = float(upper_candidates.max())
    box_half = category_width * 0.19
    cap_half = category_width * 0.16
    _add_line(slide, center_x, _role_y(lower, top), center_x, _role_y(upper, top), color=line_color, width=0.8, name=f"{name}_Whisker")
    _add_line(slide, center_x - cap_half, _role_y(lower, top), center_x + cap_half, _role_y(lower, top), color=line_color, width=0.8, name=f"{name}_LowerCap")
    _add_line(slide, center_x - cap_half, _role_y(upper, top), center_x + cap_half, _role_y(upper, top), color=line_color, width=0.8, name=f"{name}_UpperCap")
    box_top = _role_y(q3, top)
    box_bottom = _role_y(q1, top)
    box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(center_x - box_half),
        Inches(box_top),
        Inches(box_half * 2),
        Inches(max(box_bottom - box_top, 0.015)),
    )
    box.name = f"{name}_IQR"
    box.fill.solid()
    box.fill.fore_color.rgb = _rgb(fill_color)
    box.line.color.rgb = _rgb(line_color)
    box.line.width = Pt(0.8)
    _add_line(
        slide,
        center_x - box_half * 1.25,
        _role_y(median, top),
        center_x + box_half * 1.25,
        _role_y(median, top),
        color=TEXT,
        width=1.15,
        name=f"{name}_Median",
    )


def _identity(row: Mapping[str, Any]) -> tuple[int, str, int]:
    return int(float(row["dataset_id"])), str(row["dataset_name"]), int(float(row["source_sample_index"]))


def _validate_rows(
    anchor_rows: Sequence[Mapping[str, Any]],
    role_rows: Sequence[Mapping[str, Any]],
    *,
    subset_order: Sequence[str],
    expected_count: int,
) -> None:
    if len(role_rows) != int(expected_count):
        raise ValueError(f"Expected {expected_count} role records, found {len(role_rows)}")
    role_ids = {_identity(row) for row in role_rows}
    if len(role_ids) != len(role_rows):
        raise ValueError("Role diagnostics contain duplicate identities.")
    if len(anchor_rows) != 2 * int(expected_count):
        raise ValueError(f"Expected {2 * expected_count} anchor rows, found {len(anchor_rows)}")
    for role in ("background", "structural"):
        selected = [row for row in anchor_rows if str(row["role"]) == role]
        if {_identity(row) for row in selected} != role_ids:
            raise ValueError(f"Anchor identities for {role} do not match role diagnostics.")
    observed = {str(row["dataset_name"]) for row in role_rows}
    if observed != set(subset_order):
        raise ValueError(f"Subset mismatch: expected {list(subset_order)}, got {sorted(observed)}")
    for row in role_rows:
        for key in ("eta_b", "eta_s", "chi_bs"):
            value = float(row[key])
            if not np.isfinite(value) or not 0.0 <= value <= 1.0:
                raise ValueError(f"{key} must be finite and in [0,1], got {value}")


def _bootstrap_delta(values_a: np.ndarray, values_b: np.ndarray, *, seed: int) -> tuple[float, float, float]:
    delta = np.asarray(values_a, dtype=np.float64) - np.asarray(values_b, dtype=np.float64)
    rng = np.random.default_rng(seed)
    means = np.asarray([delta[rng.integers(0, delta.size, size=delta.size)].mean() for _ in range(1000)])
    return float(delta.mean()), float(np.quantile(means, 0.025)), float(np.quantile(means, 0.975))


def export_condition_contract_pptx(
    *,
    anchor_rows: Sequence[Mapping[str, Any]],
    role_rows: Sequence[Mapping[str, Any]],
    output_path: str | Path,
    metadata_output_path: str | Path | None = None,
    subset_order: Sequence[str] = SUBSET_ORDER,
    expected_count: int = 33_600,
) -> Path:
    """Create a one-slide native-shape deck for the compact Figure 3."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches

    subset_order = tuple(str(name) for name in subset_order)
    _validate_rows(anchor_rows, role_rows, subset_order=subset_order, expected_count=expected_count)
    output_path = _resolve_path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    anchors: dict[str, tuple[np.ndarray, np.ndarray]] = {}
    for role in ("background", "structural"):
        selected = [row for row in anchor_rows if str(row["role"]) == role]
        anchors[role] = (
            np.asarray([float(row["matched"]) for row in selected]),
            np.asarray([float(row["shuffled"]) for row in selected]),
        )

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb("FFFFFF")

    _add_text(slide, "(a) Physical anchor calibration", left=0.48, top=0.08, width=4.70, height=0.36, size=18, bold=True, align="center", name="PanelA_Title")
    for tick in (-0.5, 0.0, 0.5, 1.0):
        y = _anchor_y(tick)
        _add_line(slide, ANCHOR_LEFT, y, ANCHOR_LEFT + ANCHOR_WIDTH, y, color=GRID, width=0.5, name=f"AnchorGrid_{tick:.1f}")
        _add_text(slide, f"{tick:.1f}", left=0.42, top=y - 0.11, width=0.32, height=0.22, size=10.5, color=AXIS, align="right", name=f"AnchorTick_{tick:.1f}")
    _add_line(slide, ANCHOR_LEFT, ANCHOR_TOP, ANCHOR_LEFT, ANCHOR_TOP + ANCHOR_HEIGHT, color=AXIS, width=0.9, name="AnchorYAxis")
    _add_line(slide, ANCHOR_LEFT, ANCHOR_TOP + ANCHOR_HEIGHT, ANCHOR_LEFT + ANCHOR_WIDTH, ANCHOR_TOP + ANCHOR_HEIGHT, color=AXIS, width=0.9, name="AnchorXAxis")
    _add_text(slide, "Cosine similarity", left=0.02, top=1.63, width=0.34, height=2.05, size=12.5, rotation=270, align="center", name="AnchorYAxisTitle")

    centers = {
        ("background", "matched"): 1.60,
        ("background", "shuffled"): 2.38,
        ("structural", "matched"): 3.33,
        ("structural", "shuffled"): 4.11,
    }
    for role, color, light in (("background", GREEN, GREEN_LIGHT), ("structural", BLUE, BLUE_LIGHT)):
        matched, shuffled = anchors[role]
        _add_violin(slide, matched, center_x=centers[(role, "matched")], max_half_width=0.27, fill_color=light, line_color=color, name=f"Violin_{role}_matched")
        _add_violin(slide, shuffled, center_x=centers[(role, "shuffled")], max_half_width=0.27, fill_color=GRAY_LIGHT, line_color=GRAY, name=f"Violin_{role}_shuffled")
        for kind, values, line_color in (("matched", matched, color), ("shuffled", shuffled, GRAY)):
            median = float(np.median(values))
            x = centers[(role, kind)]
            _add_line(slide, x - 0.22, _anchor_y(median), x + 0.22, _anchor_y(median), color=TEXT, width=1.35, name=f"Median_{role}_{kind}")
        delta, lower, upper = _bootstrap_delta(matched, shuffled, seed=17 if role == "background" else 29)
        _add_text(
            slide,
            f"{role[0].upper()} med={np.median(matched):.2f}; Δ={delta:.2f} [{lower:.2f},{upper:.2f}]",
            left=1.08 if role == "background" else 2.94,
            top=0.45,
            width=1.83,
            height=0.25,
            size=9.2,
            color=MUTED,
            align="center",
            name=f"AnchorSummary_{role}",
        )
    _add_text(slide, "background", left=1.23, top=4.67, width=1.53, height=0.28, size=12.0, align="center", name="AnchorGroup_background")
    _add_text(slide, "structure", left=2.95, top=4.67, width=1.55, height=0.28, size=12.0, align="center", name="AnchorGroup_structural")
    _add_circle(slide, left=1.22, top=5.12, diameter=0.17, color=GREEN, name="LegendMatchedDot")
    _add_text(slide, "matched", left=1.43, top=5.06, width=0.77, height=0.27, size=11.0, name="LegendMatchedText")
    _add_circle(slide, left=2.35, top=5.12, diameter=0.17, color=GRAY, name="LegendShuffledDot")
    _add_text(slide, "shuffled", left=2.56, top=5.06, width=0.82, height=0.27, size=11.0, name="LegendShuffledText")

    _add_text(slide, "(b) Role specialization diagnostics", left=5.28, top=0.08, width=7.70, height=0.36, size=18, bold=True, align="center", name="PanelB_Title")
    metric_specs = (
        ("eta_b", "η_B ↑", GREEN, GREEN_LIGHT),
        ("eta_s", "η_S ↑", BLUE, BLUE_LIGHT),
        ("chi_bs", "χ_BS ↓", PURPLE, PURPLE_LIGHT),
    )
    category_width = ROLE_WIDTH / len(subset_order)
    for metric_index, (key, label, color, fill) in enumerate(metric_specs):
        top = ROLE_TOPS[metric_index]
        for tick in (0.0, 0.5, 1.0):
            y = _role_y(tick, top)
            _add_line(slide, ROLE_LEFT, y, ROLE_LEFT + ROLE_WIDTH, y, color=GRID, width=0.5, name=f"RoleGrid_{key}_{tick:.1f}")
            _add_text(slide, f"{tick:.1f}", left=5.27, top=y - 0.10, width=0.35, height=0.20, size=9.5, color=AXIS, align="right", name=f"RoleTick_{key}_{tick:.1f}")
        _add_line(slide, ROLE_LEFT, top, ROLE_LEFT, top + ROLE_HEIGHT, color=AXIS, width=0.8, name=f"RoleYAxis_{key}")
        _add_line(slide, ROLE_LEFT, top + ROLE_HEIGHT, ROLE_LEFT + ROLE_WIDTH, top + ROLE_HEIGHT, color=AXIS, width=0.8, name=f"RoleXAxis_{key}")
        _add_text(slide, label, left=5.03, top=top + 0.26, width=0.52, height=0.38, size=12.5, color=color, italic=True, align="center", name=f"RoleMetric_{key}")
        values_all = np.asarray([float(row[key]) for row in role_rows])
        global_median = float(np.median(values_all))
        _add_line(slide, ROLE_LEFT, _role_y(global_median, top), ROLE_LEFT + ROLE_WIDTH, _role_y(global_median, top), color=color, width=1.0, dashed=True, name=f"RoleGlobalMedian_{key}")
        _add_text(slide, f"global med. {global_median:.2f}", left=11.08, top=top + 0.02, width=1.48, height=0.24, size=9.7, color=color, align="right", name=f"RoleGlobalMedianLabel_{key}")
        for subset_index, subset_name in enumerate(subset_order):
            values = np.asarray([float(row[key]) for row in role_rows if str(row["dataset_name"]) == subset_name])
            center_x = ROLE_LEFT + category_width * (subset_index + 0.5)
            _add_role_boxplot(
                slide,
                values,
                center_x=center_x,
                top=top,
                category_width=category_width,
                line_color=color,
                fill_color=fill,
                name=f"Role_{key}_{subset_name}",
            )
            if metric_index == len(metric_specs) - 1:
                _add_text(
                    slide,
                    subset_name,
                    left=center_x - category_width * 0.53,
                    top=top + ROLE_HEIGHT + 0.03,
                    width=category_width * 1.05,
                    height=0.56,
                    size=8.6,
                    align="center",
                    rotation=315,
                    name=f"RoleSubsetLabel_{subset_name}",
                )

    presentation.save(output_path)
    reopened = Presentation(output_path)
    shapes = list(reopened.slides[0].shapes)
    picture_count = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes)
    freeform_count = sum(shape.shape_type == MSO_SHAPE_TYPE.FREEFORM for shape in shapes)
    if picture_count != 0 or freeform_count != 4:
        raise RuntimeError(f"PPTX editability check failed: pictures={picture_count}, freeforms={freeform_count}")

    if metadata_output_path is not None:
        metadata_path = _resolve_path(metadata_output_path)
        metadata_path.parent.mkdir(parents=True, exist_ok=True)
        metadata_path.write_text(
            json.dumps(
                {
                    "presentation": str(output_path),
                    "source_anchor_rows": len(anchor_rows),
                    "source_role_rows": len(role_rows),
                    "subset_order": list(subset_order),
                    "slide_size_inches": [SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN],
                    "shape_count": len(shapes),
                    "freeform_violin_count": freeform_count,
                    "picture_count": picture_count,
                    "all_elements_native_editable": picture_count == 0,
                    "main_figure_panels": ["anchor_calibration", "role_specialization"],
                    "font": FONT_NAME,
                },
                indent=2,
                sort_keys=True,
            ),
            encoding="utf-8",
        )
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Export compact Figure 3 as editable native PowerPoint shapes.")
    parser.add_argument("--anchor-csv", default=str(DEFAULT_OUTPUT_DIR / "anchor_calibration.csv"))
    parser.add_argument("--role-csv", default=str(DEFAULT_OUTPUT_DIR / "role_separation.csv"))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT_DIR / "figure3_condition_learning_editable.pptx"))
    parser.add_argument("--output-metadata", default=str(DEFAULT_OUTPUT_DIR / "figure3_condition_learning_pptx_metadata.json"))
    parser.add_argument("--expected-count", type=int, default=33_600)
    args = parser.parse_args()
    output = export_condition_contract_pptx(
        anchor_rows=_read_csv(args.anchor_csv),
        role_rows=_read_csv(args.role_csv),
        output_path=args.output,
        metadata_output_path=args.output_metadata,
        subset_order=SUBSET_ORDER,
        expected_count=args.expected_count,
    )
    print(json.dumps({"editable_pptx": str(output)}, indent=2))


if __name__ == "__main__":
    main()
