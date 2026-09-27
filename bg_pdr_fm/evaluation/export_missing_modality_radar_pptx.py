"""Export the missing-modality radar figure as editable PowerPoint shapes."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
from typing import Any, Mapping, Sequence

import numpy as np

from bg_pdr_fm.evaluation.visualize_missing_modality_radar import (
    METHOD_STYLES,
    RADAR_DATASET_LABELS,
    RADAR_DATASET_ORDER,
    RADAR_METHOD_ORDER,
    RADAR_MODE_LABELS,
    RADAR_MODE_ORDER,
    RADAR_PANEL_RANGES,
    RADAR_PANEL_TICKS,
    load_radar_rows,
    validate_radar_rows,
)


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_DIR = REPO_ROOT / "docs" / "paper" / "AAAI2027" / "figures" / "missing_modality_radar"
DEFAULT_OUTPUT = DEFAULT_OUTPUT_DIR / "missing_modality_radar_4panel_editable.pptx"
DEFAULT_METADATA = DEFAULT_OUTPUT_DIR / "missing_modality_radar_4panel_pptx_metadata.json"

SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 5.00
PANEL_CENTER_X = (1.65, 4.99, 8.33, 11.67)
PANEL_CENTER_Y = 2.05
PANEL_RADIUS = 1.28
FONT_NAME = "Times New Roman"

GRID_COLOR = "B8B8B8"
AXIS_COLOR = "666666"
TEXT_COLOR = "252525"
MUTED_COLOR = "666666"
WHITE = "FFFFFF"

PPT_METHOD_STYLES = {
    "InversionNet": {"dash": "solid", "marker": "oval", "rotation": 0.0},
    "VelocityGAN": {"dash": "dash", "marker": "rectangle", "rotation": 0.0},
    "UPFWI": {"dash": "dash_dot", "marker": "triangle", "rotation": 0.0},
    "Auto-Linear": {"dash": "dot", "marker": "diamond", "rotation": 0.0},
    "Latent U-Net (Large)": {"dash": "long_dash", "marker": "triangle", "rotation": 180.0},
    "PD-BG-RFM": {"dash": "solid", "marker": "star", "rotation": 0.0},
}


def _resolve_path(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else REPO_ROOT / path


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _rgb(hex_value: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_value.lstrip("#"))


def _add_text(
    slide: Any,
    text: str,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    size: float,
    color: str = TEXT_COLOR,
    bold: bool = False,
    italic: bool = False,
    align: str = "left",
    name: str,
) -> Any:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    textbox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    textbox.name = name
    textbox.text_frame.clear()
    textbox.text_frame.margin_left = 0
    textbox.text_frame.margin_right = 0
    textbox.text_frame.margin_top = 0
    textbox.text_frame.margin_bottom = 0
    textbox.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = textbox.text_frame.paragraphs[0]
    paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)
    return textbox


def _add_line(
    slide: Any,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str,
    width: float,
    dash: str = "solid",
    name: str,
) -> Any:
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches, Pt

    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.name = name
    line.line.color.rgb = _rgb(color)
    line.line.width = Pt(width)
    dash_styles = {
        "solid": MSO_LINE_DASH_STYLE.SOLID,
        "dash": MSO_LINE_DASH_STYLE.DASH,
        "dash_dot": MSO_LINE_DASH_STYLE.DASH_DOT,
        "dot": MSO_LINE_DASH_STYLE.ROUND_DOT,
        "long_dash": MSO_LINE_DASH_STYLE.LONG_DASH,
    }
    line.line.dash_style = dash_styles[dash]
    return line


def _add_circle(slide: Any, *, center_x: float, center_y: float, radius: float, name: str) -> Any:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(center_x - radius),
        Inches(center_y - radius),
        Inches(2.0 * radius),
        Inches(2.0 * radius),
    )
    circle.name = name
    circle.fill.background()
    circle.line.color.rgb = _rgb(GRID_COLOR)
    circle.line.width = Pt(0.55)
    return circle


def _add_marker(
    slide: Any,
    *,
    x: float,
    y: float,
    marker: str,
    color: str,
    rotation: float,
    size: float,
    name: str,
) -> Any:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    marker_shapes = {
        "oval": MSO_SHAPE.OVAL,
        "rectangle": MSO_SHAPE.RECTANGLE,
        "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
        "diamond": MSO_SHAPE.DIAMOND,
        "star": MSO_SHAPE.STAR_5_POINT,
    }
    shape = slide.shapes.add_shape(
        marker_shapes[marker],
        Inches(x - size / 2.0),
        Inches(y - size / 2.0),
        Inches(size),
        Inches(size),
    )
    shape.name = name
    shape.rotation = rotation
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(WHITE)
    shape.line.color.rgb = _rgb(color)
    shape.line.width = Pt(0.55)
    return shape


def _polar_point(center_x: float, center_y: float, radius: float, theta: float) -> tuple[float, float]:
    return (
        center_x + radius * math.sin(theta),
        center_y - radius * math.cos(theta),
    )


def _value_radius(value: float, mode: str) -> float:
    lower, upper = RADAR_PANEL_RANGES[mode]
    if not lower <= value <= upper:
        raise ValueError(f"{mode} SSIM value {value} is outside editable display range [{lower}, {upper}]")
    return PANEL_RADIUS * (value - lower) / (upper - lower)


def _matrix(rows: Sequence[Mapping[str, Any]]) -> dict[str, dict[str, list[float]]]:
    methods = tuple(RADAR_METHOD_ORDER)
    values = {
        mode: {method: [0.0] * len(RADAR_DATASET_ORDER) for method in methods}
        for mode in RADAR_MODE_ORDER
    }
    dataset_index = {name: index for index, name in enumerate(RADAR_DATASET_ORDER)}
    for row in rows:
        values[str(row["missing_mode"])][str(row["method"])][dataset_index[str(row["dataset_name"])]] = float(
            row["ssim"]
        )
    return values


def _add_radar_panel(
    slide: Any,
    *,
    mode: str,
    panel_index: int,
    values: Mapping[str, Sequence[float]],
) -> list[Any]:
    center_x = PANEL_CENTER_X[panel_index]
    center_y = PANEL_CENTER_Y
    ticks = RADAR_PANEL_TICKS[mode]
    lower, upper = RADAR_PANEL_RANGES[mode]
    theta = np.linspace(0.0, 2.0 * math.pi, len(RADAR_DATASET_ORDER), endpoint=False)
    shapes: list[Any] = []

    shapes.append(
        _add_text(
            slide,
            f"({chr(ord('a') + panel_index)}) {RADAR_MODE_LABELS[mode]}",
            left=center_x - 1.38,
            top=0.12,
            width=2.76,
            height=0.34,
            size=13.5,
            bold=True,
            align="center",
            name=f"RadarPanel_{RADAR_MODE_LABELS[mode]}",
        )
    )
    for tick in ticks:
        normalized = (tick - lower) / (upper - lower)
        if normalized > 1e-8:
            shapes.append(
                _add_circle(
                    slide,
                    center_x=center_x,
                    center_y=center_y,
                    radius=PANEL_RADIUS * normalized,
                    name=f"Grid_{mode}_{tick:.2f}",
                )
            )
        tick_x, tick_y = _polar_point(center_x, center_y, max(PANEL_RADIUS * normalized, 0.03), math.pi / 4.0)
        shapes.append(
            _add_text(
                slide,
                f"{tick:.2f}",
                left=tick_x + 0.02,
                top=tick_y - 0.07,
                width=0.38,
                height=0.16,
                size=6.2,
                color=MUTED_COLOR,
                name=f"RadialTick_{mode}_{tick:.2f}",
            )
        )

    for axis_index, angle in enumerate(theta):
        end_x, end_y = _polar_point(center_x, center_y, PANEL_RADIUS, angle)
        shapes.append(
            _add_line(
                slide,
                center_x,
                center_y,
                end_x,
                end_y,
                color=GRID_COLOR,
                width=0.55,
                name=f"Spoke_{mode}_{axis_index}",
            )
        )
        label_x, label_y = _polar_point(center_x, center_y, PANEL_RADIUS + 0.20, angle)
        shapes.append(
            _add_text(
                slide,
                RADAR_DATASET_LABELS[RADAR_DATASET_ORDER[axis_index]],
                left=label_x - 0.27,
                top=label_y - 0.08,
                width=0.54,
                height=0.17,
                size=7.3,
                align="center",
                name=f"AxisLabel_{mode}_{RADAR_DATASET_ORDER[axis_index]}",
            )
        )

    for method, method_values in values.items():
        color = str(METHOD_STYLES[method]["color"]).lstrip("#")
        style = PPT_METHOD_STYLES[method]
        panel_label = RADAR_MODE_LABELS[mode]
        line_width = 1.9 if method == "PD-BG-RFM" else 0.95
        marker_size = 0.095 if method == "PD-BG-RFM" else 0.068
        points = [
            _polar_point(center_x, center_y, _value_radius(float(value), mode), angle)
            for value, angle in zip(method_values, theta)
        ]
        for segment_index in range(len(points)):
            x1, y1 = points[segment_index]
            x2, y2 = points[(segment_index + 1) % len(points)]
            shapes.append(
                _add_line(
                    slide,
                    x1,
                    y1,
                    x2,
                    y2,
                    color=color,
                    width=line_width,
                    dash=str(style["dash"]),
                    name=f"RadarCurve_{method}_{panel_label}_segment_{segment_index}",
                )
            )
        for marker_index, (x, y) in enumerate(points):
            shapes.append(
                _add_marker(
                    slide,
                    x=x,
                    y=y,
                    marker=str(style["marker"]),
                    color=color,
                    rotation=float(style["rotation"]),
                    size=marker_size,
                    name=f"RadarMarker_{method}_{panel_label}_{marker_index}",
                )
            )
    return shapes


def _add_legend(slide: Any) -> list[Any]:
    from pptx.util import Inches

    shapes: list[Any] = []
    legend_columns = (
        ("InversionNet", "VelocityGAN"),
        ("UPFWI", "Auto-Linear"),
        ("Latent U-Net (Large)", "PD-BG-RFM"),
    )
    for column, methods in enumerate(legend_columns):
        x = 2.00 + column * 3.05
        for row, method in enumerate(methods):
            y = 4.20 + row * 0.29
            color = str(METHOD_STYLES[method]["color"]).lstrip("#")
            style = PPT_METHOD_STYLES[method]
            line_width = 1.9 if method == "PD-BG-RFM" else 0.95
            _add_line(
                slide,
                x,
                y + 0.075,
                x + 0.30,
                y + 0.075,
                color=color,
                width=line_width,
                dash=str(style["dash"]),
                name=f"LegendLine_{method}",
            )
            _add_marker(
                slide,
                x=x + 0.15,
                y=y + 0.075,
                marker=str(style["marker"]),
                color=color,
                rotation=float(style["rotation"]),
                size=0.075 if method != "PD-BG-RFM" else 0.10,
                name=f"Legend_{method}",
            )
            _add_text(
                slide,
                method,
                left=x + 0.42,
                top=y,
                width=2.25,
                height=0.18,
                size=7.4,
                name=f"LegendText_{method}",
            )
    return shapes


def export_missing_modality_radar_pptx(
    *,
    rows: Sequence[Mapping[str, Any]],
    output_path: str | Path,
    metadata_output_path: str | Path | None = None,
    source_csv: str | Path | None = None,
) -> Path:
    """Create a one-slide, native-shape PowerPoint for the radar figure."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches

    validate_radar_rows(rows)
    matrix = _matrix(rows)
    for mode in RADAR_MODE_ORDER:
        for method_values in matrix[mode].values():
            for value in method_values:
                _value_radius(value, mode)

    output = _resolve_path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(WHITE)

    shape_list: list[Any] = []
    for panel_index, mode in enumerate(RADAR_MODE_ORDER):
        shape_list.extend(
            _add_radar_panel(
                slide,
                mode=mode,
                panel_index=panel_index,
                values=matrix[mode],
            )
        )
    shape_list.extend(_add_legend(slide))
    shape_list.append(
        _add_text(
            slide,
            "Raw SSIM; panel-specific radial display ranges",
            left=9.65,
            top=4.74,
            width=3.20,
            height=0.16,
            size=5.8,
            color=MUTED_COLOR,
            italic=True,
            align="right",
            name="ScaleNote",
        )
    )

    presentation.save(output)
    reopened = Presentation(output)
    shapes = list(reopened.slides[0].shapes)
    picture_count = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes)
    line_count = sum(shape.shape_type == MSO_SHAPE_TYPE.LINE for shape in shapes)
    if picture_count != 0 or line_count <= 100:
        raise RuntimeError(f"PPTX editability check failed: pictures={picture_count}, lines={line_count}")

    if metadata_output_path is not None:
        metadata = _resolve_path(metadata_output_path)
        metadata.parent.mkdir(parents=True, exist_ok=True)
        source = _resolve_path(source_csv) if source_csv is not None else None
        payload = {
            "presentation": str(output),
            "source_csv": str(source) if source is not None else "",
            "source_csv_sha256": _sha256(source) if source is not None and source.is_file() else "",
            "methods": list(RADAR_METHOD_ORDER),
            "datasets": list(RADAR_DATASET_ORDER),
            "modes": list(RADAR_MODE_ORDER),
            "num_rows": len(rows),
            "slide_size_inches": [SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN],
            "panel_radial_ranges": {mode: list(RADAR_PANEL_RANGES[mode]) for mode in RADAR_MODE_ORDER},
            "panel_radial_ticks": {mode: list(RADAR_PANEL_TICKS[mode]) for mode in RADAR_MODE_ORDER},
            "shape_count": len(shapes),
            "line_count": line_count,
            "picture_count": picture_count,
            "all_elements_native_editable": picture_count == 0,
            "contains_embedded_media": picture_count != 0,
            "font": FONT_NAME,
        }
        metadata.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-csv", default=str(DEFAULT_OUTPUT_DIR / "missing_modality_radar_source.csv"))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--output-metadata", default=str(DEFAULT_METADATA))
    args = parser.parse_args()
    source = _resolve_path(args.input_csv)
    output = export_missing_modality_radar_pptx(
        rows=load_radar_rows(source),
        output_path=args.output,
        metadata_output_path=args.output_metadata,
        source_csv=source,
    )
    print(json.dumps({"editable_pptx": str(output)}, indent=2))


if __name__ == "__main__":
    main()
