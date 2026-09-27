"""Export the Figure 1 motivation statistics as editable PowerPoint shapes.

The source statistics are the frozen CSV/NPZ assets used by the PDF renderer.
No PDF or raster image is embedded in the resulting presentation.
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any, Mapping, Sequence

import numpy as np


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT_DIR = REPO_ROOT / "docs" / "paper" / "AAAI2027" / "figures" / "figure1_motivation"
DEFAULT_OUTPUT = DEFAULT_INPUT_DIR / "motivation_statistics_combined_v2_editable.pptx"
DEFAULT_METADATA_OUTPUT = DEFAULT_INPUT_DIR / "motivation_statistics_combined_v2_pptx_metadata.json"

SOURCE_WIDTH_PT = 504.0
SOURCE_HEIGHT_PT = 194.4
SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = SLIDE_WIDTH_IN * SOURCE_HEIGHT_PT / SOURCE_WIDTH_PT

FONT_NAME = "Times New Roman"
TEXT = "252525"
MUTED = "596270"
AXIS = "666666"
GRID = "D9D9D9"
PAIR = "B5B5B5"
BACKGROUND = "2A9D6F"
STRUCTURE = "D55E00"

MODALITIES = ("migrated_image", "horizon", "rms_vel", "well_log_mask")
MODALITY_LABELS = {
    "migrated_image": "PoSTM",
    "horizon": "Horizon",
    "rms_vel": "RMS velocity",
    "well_log_mask": "Well log + mask",
}
COMPONENTS = ("background", "structure")
SUBSETS = (
    "FlatVelA",
    "FlatVelB",
    "CurveVelA",
    "CurveVelB",
    "FlatFaultA",
    "FlatFaultB",
    "CurveFaultA",
    "CurveFaultB",
)


def _resolve_path(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else REPO_ROOT / path


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _load_data(input_dir: str | Path) -> dict[str, Any]:
    root = _resolve_path(input_dir)
    required = (
        "alignment_by_subset.csv",
        "effective_rank_by_subset.csv",
        "explained_variance_curves.npz",
        "analysis_config.json",
    )
    missing = [name for name in required if not (root / name).is_file()]
    if missing:
        raise FileNotFoundError(f"Missing motivation artifacts under {root}: {missing}")

    alignment_rows = _read_csv(root / "alignment_by_subset.csv")
    rank_rows = _read_csv(root / "effective_rank_by_subset.csv")
    with np.load(root / "explained_variance_curves.npz") as archive:
        background_curves = np.asarray(archive["background_curves"], dtype=np.float64).copy()
        structure_curves = np.asarray(archive["structure_curves"], dtype=np.float64).copy()
    config = json.loads((root / "analysis_config.json").read_text(encoding="utf-8"))

    alignment = {
        (str(row["modality"]), str(row["component"])): float(row["rbf_cka_median"])
        for row in alignment_rows
        if row.get("scope") == "all_subsets_summary"
    }
    expected_alignment = {(modality, component) for modality in MODALITIES for component in COMPONENTS}
    if set(alignment) != expected_alignment:
        raise ValueError("alignment_by_subset.csv does not contain the expected 4x2 summary matrix")

    rank = {
        (str(row["subset"]), str(row["component"])): float(row["effective_rank"])
        for row in rank_rows
        if row.get("scope") == "subset"
    }
    expected_rank = {(subset, component) for subset in SUBSETS for component in COMPONENTS}
    if set(rank) != expected_rank:
        raise ValueError("effective_rank_by_subset.csv does not contain the expected 8x2 subset matrix")
    if background_curves.ndim != 2 or structure_curves.shape != background_curves.shape:
        raise ValueError("Explained-variance curves must be two arrays with identical 2D shapes")
    if background_curves.shape[0] != len(SUBSETS):
        raise ValueError(f"Expected {len(SUBSETS)} explained-variance curves")

    return {
        "input_dir": root,
        "alignment_matrix": np.asarray(
            [[alignment[(modality, component)] for component in COMPONENTS] for modality in MODALITIES],
            dtype=np.float64,
        ),
        "rank_pairs": np.asarray(
            [[rank[(subset, component)] for component in COMPONENTS] for subset in SUBSETS],
            dtype=np.float64,
        ),
        "background_curves": background_curves,
        "structure_curves": structure_curves,
        "config": config,
    }


def _rgb(value: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(value)


def _add_text(
    slide: Any,
    text: str,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    size: float,
    color: str = TEXT,
    bold: bool = False,
    italic: bool = False,
    align: str = "left",
    rotation: float = 0.0,
    name: str | None = None,
) -> Any:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if name:
        shape.name = name
    shape.rotation = rotation
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)
    return shape


def _add_line(
    slide: Any,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str = AXIS,
    width: float = 0.6,
    name: str | None = None,
) -> Any:
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if name:
        shape.name = name
    shape.line.color.rgb = _rgb(color)
    shape.line.width = Pt(width)
    return shape


def _add_rect(
    slide: Any,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    fill: str,
    line: str | None = None,
    line_width: float = 0.4,
    name: str | None = None,
) -> Any:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.color.rgb = _rgb(line or fill)
    shape.line.width = Pt(line_width)
    return shape


def _add_circle(
    slide: Any,
    *,
    center_x: float,
    center_y: float,
    diameter: float,
    fill: str,
    line: str = "FFFFFF",
    name: str,
) -> Any:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(center_x - diameter / 2),
        Inches(center_y - diameter / 2),
        Inches(diameter),
        Inches(diameter),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.color.rgb = _rgb(line)
    shape.line.width = Pt(0.35)
    return shape


def _add_freeform(
    slide: Any,
    points: Sequence[tuple[float, float]],
    *,
    fill: str | None,
    line: str,
    line_width: float,
    name: str,
    close: bool = False,
) -> Any:
    from pptx.util import Inches, Pt

    if len(points) < 2:
        raise ValueError("A freeform requires at least two points")
    local_points = [(int(round(x * 1000)), int(round(y * 1000))) for x, y in points]
    scale = float(Inches(1)) / 1000.0
    builder = slide.shapes.build_freeform(
        start_x=local_points[0][0],
        start_y=local_points[0][1],
        scale=(scale, scale),
    )
    builder.add_line_segments(local_points[1:], close=close)
    shape = builder.convert_to_shape()
    shape.name = name
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.color.rgb = _rgb(line)
    shape.line.width = Pt(line_width)
    return shape


def _sample_cividis(value: float) -> str:
    from matplotlib import colormaps

    red, green, blue, _ = colormaps["cividis"](float(np.clip(value, 0.0, 1.0)))
    return "{:02X}{:02X}{:02X}".format(round(red * 255), round(green * 255), round(blue * 255))


def _draw_alignment_panel(slide: Any, matrix: np.ndarray) -> int:
    left, top = 1.10, 1.04
    width, height = 2.55, 3.36
    cell_w, cell_h = width / 2.0, height / 4.0
    shape_count = 0
    _add_text(
        slide,
        "(a) Modality-component alignment",
        left=0.36,
        top=0.27,
        width=4.15,
        height=0.30,
        size=12.0,
        name="Panel_A_Title",
    )
    for row_index, modality in enumerate(MODALITIES):
        y = top + row_index * cell_h
        _add_text(
            slide,
            MODALITY_LABELS[modality],
            left=0.05,
            top=y,
            width=0.98,
            height=cell_h,
            size=8.2,
            align="right",
            name=f"Alignment_RowLabel_{modality}",
        )
        for column_index in range(2):
            x = left + column_index * cell_w
            value = float(matrix[row_index, column_index])
            rect = _add_rect(
                slide,
                left=x,
                top=y,
                width=cell_w,
                height=cell_h,
                fill=_sample_cividis(value),
                line="FFFFFF",
                line_width=1.0,
                name=f"Alignment_Cell_{row_index}_{column_index}",
            )
            del rect
            shape_count += 1
            luminance = 0.2126 * int(_sample_cividis(value)[0:2], 16) / 255.0
            luminance += 0.7152 * int(_sample_cividis(value)[2:4], 16) / 255.0
            luminance += 0.0722 * int(_sample_cividis(value)[4:6], 16) / 255.0
            _add_text(
                slide,
                f"{value:.2f}",
                left=x,
                top=y,
                width=cell_w,
                height=cell_h,
                size=9.0,
                color="FFFFFF" if luminance < 0.50 else TEXT,
                align="center",
                name=f"Alignment_Value_{row_index}_{column_index}",
            )
            shape_count += 1
    _add_text(slide, "Background B", left=left, top=top + height + 0.05, width=cell_w, height=0.28, size=8.0, align="center", name="Alignment_ColumnLabel_B")
    _add_text(slide, "Structure S", left=left + cell_w, top=top + height + 0.05, width=cell_w, height=0.28, size=8.0, align="center", name="Alignment_ColumnLabel_S")
    shape_count += 2

    bar_left, bar_top, bar_width, bar_height = 3.88, top, 0.14, height
    steps = 32
    for index in range(steps):
        y = bar_top + bar_height * index / steps
        _add_rect(
            slide,
            left=bar_left,
            top=y,
            width=bar_width,
            height=bar_height / steps + 0.002,
            fill=_sample_cividis(1.0 - index / (steps - 1)),
            line=_sample_cividis(1.0 - index / (steps - 1)),
            line_width=0.1,
            name=f"Alignment_Colorbar_{index:02d}",
        )
        shape_count += 1
    _add_rect(slide, left=bar_left, top=bar_top, width=bar_width, height=bar_height, fill="FFFFFF", line=AXIS, line_width=0.4, name="Alignment_Colorbar_Outline")
    shape_count += 1
    for tick in (0.0, 0.2, 0.4, 0.6, 0.8, 1.0):
        y = bar_top + bar_height * (1.0 - tick)
        _add_line(slide, bar_left + bar_width, y, bar_left + bar_width + 0.06, y, color=AXIS, width=0.45, name=f"Alignment_Colorbar_Tick_{tick:.1f}")
        _add_text(slide, f"{tick:.1f}", left=bar_left + bar_width + 0.08, top=y - 0.10, width=0.28, height=0.20, size=6.3, color=MUTED, name=f"Alignment_Colorbar_Label_{tick:.1f}")
        shape_count += 3
    _add_text(slide, "Median RBF CKA", left=bar_left - 0.18, top=top + height / 2 - 0.55, width=0.16, height=1.10, size=6.4, color=MUTED, rotation=90, align="center", name="Alignment_Colorbar_Title")
    shape_count += 1
    return shape_count


def _draw_axes(slide: Any, *, left: float, top: float, width: float, height: float, y_ticks: Sequence[float], x_ticks: Sequence[float], x_to: Any, y_to: Any, prefix: str) -> int:
    count = 0
    for tick in y_ticks:
        y = y_to(float(tick))
        _add_line(slide, left, y, left + width, y, color=GRID, width=0.45, name=f"{prefix}_YGrid_{tick}")
        _add_text(
            slide,
            str(tick),
            left=left - 0.34,
            top=y - 0.10,
            width=0.28,
            height=0.20,
            size=6.5,
            color=MUTED,
            align="right",
            name=f"{prefix}_YLabel_{tick}",
        )
        count += 2
    for tick in x_ticks:
        x = x_to(float(tick))
        _add_line(slide, x, top + height, x, top + height + 0.05, color=AXIS, width=0.55, name=f"{prefix}_XTick_{tick}")
        _add_text(
            slide,
            str(tick),
            left=x - 0.22,
            top=top + height + 0.06,
            width=0.44,
            height=0.20,
            size=6.5,
            color=MUTED,
            align="center",
            name=f"{prefix}_XLabel_{tick}",
        )
        count += 2
    _add_line(slide, left, top + height, left + width, top + height, color=AXIS, width=0.65, name=f"{prefix}_XAxis")
    _add_line(slide, left, top, left, top + height, color=AXIS, width=0.65, name=f"{prefix}_YAxis")
    return count + 2


def _draw_effective_panel(slide: Any, background_curves: np.ndarray, structure_curves: np.ndarray, rank_pairs: np.ndarray) -> int:
    shape_count = 0
    _add_text(slide, "(b) Component effective dimension", left=5.00, top=0.27, width=5.30, height=0.30, size=12.0, name="Panel_B_Title")

    left, top, width, height = 5.43, 1.04, 4.70, 3.36
    x = np.arange(1, background_curves.shape[1] + 1, dtype=np.float64)
    x_log_max = max(3.0, float(np.log10(x[-1])))
    x_to = lambda value: left + (np.log10(max(value, 1.0)) / x_log_max) * width
    y_to = lambda value: top + (1.0 - float(value)) * height
    shape_count += _draw_axes(slide, left=left, top=top, width=width, height=height, y_ticks=(0.0, 0.2, 0.4, 0.6, 0.8, 1.0), x_ticks=(1, 10, 100, 1000), x_to=x_to, y_to=y_to, prefix="Effective")
    _add_text(slide, "Principal components", left=left + 1.45, top=top + height + 0.34, width=1.75, height=0.22, size=7.0, color=MUTED, align="center", name="Effective_XTitle")
    _add_text(slide, "Cumulative explained variance", left=left - 0.61, top=top + 0.76, width=0.22, height=1.80, size=7.0, color=MUTED, rotation=90, align="center", name="Effective_YTitle")
    shape_count += 2

    for component_curves, color, label, prefix in (
        (background_curves, BACKGROUND, "Background B", "Background"),
        (structure_curves, STRUCTURE, "Structure S", "Structure"),
    ):
        median = np.nanmedian(component_curves, axis=0)
        q1, q3 = np.nanquantile(component_curves, [0.25, 0.75], axis=0)
        band_points = [(x_to(float(xx)), y_to(float(yy))) for xx, yy in zip(x, q1)]
        band_points.extend((x_to(float(xx)), y_to(float(yy))) for xx, yy in zip(x[::-1], q3[::-1]))
        _add_freeform(slide, band_points, fill=color, line=color, line_width=0.1, name=f"{prefix}_IQR_Band", close=True)
        line_points = [(x_to(float(xx)), y_to(float(yy))) for xx, yy in zip(x, median)]
        _add_freeform(slide, line_points, fill=None, line=color, line_width=1.5, name=f"{prefix}_MedianCurve")
        shape_count += 2
        legend_x = left + (0.10 if prefix == "Background" else 1.38)
        _add_line(slide, legend_x, top + 0.14, legend_x + 0.28, top + 0.14, color=color, width=1.6, name=f"Legend_{prefix}_Line")
        _add_text(slide, label, left=legend_x + 0.34, top=top + 0.02, width=1.00, height=0.24, size=7.2, name=f"Legend_{prefix}_Label")
        shape_count += 2

    rank_left, rank_top, rank_width, rank_height = 10.78, 1.04, 2.05, 3.36
    rank_y_min, rank_y_max = 1.0, 550.0
    rank_y = lambda value: rank_top + (np.log10(rank_y_max) - np.log10(max(value, rank_y_min))) / (np.log10(rank_y_max) - np.log10(rank_y_min)) * rank_height
    rank_x = {"background": rank_left + 0.48, "structure": rank_left + 1.48}
    for tick in (1, 10, 100, 500):
        y = rank_y(tick)
        _add_line(slide, rank_left, y, rank_left + rank_width, y, color=GRID, width=0.45, name=f"Rank_YGrid_{tick}")
        _add_text(
            slide,
            str(tick),
            left=rank_left - 0.28,
            top=y - 0.10,
            width=0.24,
            height=0.20,
            size=6.5,
            color=MUTED,
            align="right",
            name=f"Rank_YLabel_{tick}",
        )
        shape_count += 2
    _add_line(slide, rank_left, rank_top + rank_height, rank_left + rank_width, rank_top + rank_height, color=AXIS, width=0.65, name="Rank_XAxis")
    _add_line(slide, rank_left, rank_top, rank_left, rank_top + rank_height, color=AXIS, width=0.65, name="Rank_YAxis")
    _add_text(slide, "Effective rank", left=rank_left, top=0.73, width=1.40, height=0.22, size=7.5, name="Rank_Title")
    _add_text(slide, "B", left=rank_x["background"] - 0.18, top=rank_top + rank_height + 0.08, width=0.36, height=0.20, size=7.2, color=BACKGROUND, align="center", name="Rank_XLabel_B")
    _add_text(slide, "S", left=rank_x["structure"] - 0.18, top=rank_top + rank_height + 0.08, width=0.36, height=0.20, size=7.2, color=STRUCTURE, align="center", name="Rank_XLabel_S")
    _add_text(slide, "r_eff", left=rank_left - 0.03, top=rank_top + 1.38, width=0.20, height=0.48, size=7.0, color=MUTED, rotation=90, align="center", name="Rank_YTitle")
    shape_count += 8
    for index, pair in enumerate(rank_pairs):
        _add_line(slide, rank_x["background"], rank_y(float(pair[0])), rank_x["structure"], rank_y(float(pair[1])), color=PAIR, width=0.65, name=f"Rank_Pair_{index}")
        _add_circle(slide, center_x=rank_x["background"], center_y=rank_y(float(pair[0])), diameter=0.11, fill=BACKGROUND, name=f"Rank_Background_{index}")
        _add_circle(slide, center_x=rank_x["structure"], center_y=rank_y(float(pair[1])), diameter=0.11, fill=STRUCTURE, name=f"Rank_Structure_{index}")
        shape_count += 3
    return shape_count


def export_motivation_statistics_pptx(
    *,
    input_dir: str | Path = DEFAULT_INPUT_DIR,
    output_path: str | Path = DEFAULT_OUTPUT,
    metadata_output_path: str | Path = DEFAULT_METADATA_OUTPUT,
) -> Path:
    """Create a one-slide, native-shape PowerPoint from frozen statistics."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches

    data = _load_data(input_dir)
    output = _resolve_path(output_path)
    metadata_output = _resolve_path(metadata_output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    metadata_output.parent.mkdir(parents=True, exist_ok=True)

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb("FFFFFF")

    shape_count = 0
    shape_count += _draw_alignment_panel(slide, data["alignment_matrix"])
    shape_count += _draw_effective_panel(
        slide,
        data["background_curves"],
        data["structure_curves"],
        data["rank_pairs"],
    )
    _add_text(
        slide,
        "Training-free statistics over eight OpenFWI subsets",
        left=0.36,
        top=4.78,
        width=5.0,
        height=0.20,
        size=6.4,
        color=MUTED,
        italic=True,
        name="Figure_Provenance_Note",
    )
    shape_count += 1

    picture_shape_count = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
    native_shape_count = len(slide.shapes)
    presentation.save(output)
    reopened = Presentation(output)
    reopened_slide = reopened.slides[0]
    reopened_picture_count = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in reopened_slide.shapes)
    metadata = {
        "source_statistics_dir": str(data["input_dir"]),
        "source_pdf_not_embedded": True,
        "slide_width_in": SLIDE_WIDTH_IN,
        "slide_height_in": SLIDE_HEIGHT_IN,
        "source_pdf_width_pt": SOURCE_WIDTH_PT,
        "source_pdf_height_pt": SOURCE_HEIGHT_PT,
        "native_shape_count": native_shape_count,
        "picture_shape_count": picture_shape_count,
        "reopened_picture_shape_count": reopened_picture_count,
        "all_elements_native_editable": picture_shape_count == 0 and reopened_picture_count == 0,
        "statistics_config": data["config"],
        "alignment_matrix": data["alignment_matrix"].tolist(),
        "rank_pairs": data["rank_pairs"].tolist(),
    }
    metadata_output.write_text(json.dumps(metadata, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="Export Figure 1 motivation statistics as editable PowerPoint shapes.")
    parser.add_argument("--input-dir", default=str(DEFAULT_INPUT_DIR))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--output-metadata", default=str(DEFAULT_METADATA_OUTPUT))
    args = parser.parse_args()
    output = export_motivation_statistics_pptx(
        input_dir=args.input_dir,
        output_path=args.output,
        metadata_output_path=args.output_metadata,
    )
    print(json.dumps({"editable_pptx": str(output)}, indent=2))


if __name__ == "__main__":
    main()
