"""Export the full-test residual transport diagnostic as editable PowerPoint shapes."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any, Mapping, Sequence

import numpy as np
from scipy.stats import gaussian_kde

from bg_pdr_fm.evaluation.visualize_residual_transport import SUBSET_ORDER


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_DIR = REPO_ROOT / "docs" / "paper" / "AAAI2027" / "figures" / "figure4_residual_transport"

SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 5.25
PLOT_LEFT_IN = 1.02
PLOT_TOP_IN = 1.08
PLOT_WIDTH_IN = 12.00
PLOT_HEIGHT_IN = 3.28
Y_MIN = 0.50
Y_MAX = 1.12

BACKGROUND_FILL = "DCEFE4"
BACKGROUND_LINE = "83BE99"
STRUCTURE_FILL = "DCE8F8"
STRUCTURE_LINE = "82ADE2"
DIAGNOSTIC_PURPLE = "8357A6"
AXIS_COLOR = "4C5563"
GRID_COLOR = "E1E5EA"
TEXT_COLOR = "252A31"
MUTED_TEXT = "596270"
FONT_NAME = "Times New Roman"


def _resolve_path(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else REPO_ROOT / path


def _read_csv(path: str | Path) -> list[dict[str, str]]:
    resolved = _resolve_path(path)
    if not resolved.is_file():
        raise FileNotFoundError(f"CSV does not exist: {resolved}")
    with resolved.open("r", newline="", encoding="utf-8") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _read_json(path: str | Path) -> dict[str, Any]:
    resolved = _resolve_path(path)
    if not resolved.is_file():
        raise FileNotFoundError(f"JSON does not exist: {resolved}")
    return json.loads(resolved.read_text(encoding="utf-8"))


def _qt_value(row: Mapping[str, Any]) -> float:
    value = row.get("q_T", row.get("q_t"))
    if value is None or str(value).strip() == "":
        raise ValueError(f"q_T is missing from row: {row}")
    result = float(value)
    if not np.isfinite(result) or result <= 0.0:
        raise ValueError(f"q_T must be finite and positive, got {result}")
    return result


def _validate_rows(rows: Sequence[Mapping[str, Any]], *, expected_count: int = 33_600) -> None:
    if len(rows) != expected_count:
        raise ValueError(f"Expected {expected_count} q_T records, found {len(rows)}")
    identities = {
        (int(float(row["dataset_id"])), str(row["dataset_name"]), int(float(row["source_sample_index"])))
        for row in rows
    }
    if len(identities) != len(rows):
        raise ValueError(f"q_T rows contain {len(rows) - len(identities)} duplicate identities")
    counts = {name: sum(str(row.get("dataset_name")) == name for row in rows) for name in SUBSET_ORDER}
    if any(count == 0 for count in counts.values()):
        raise ValueError(f"q_T rows are missing an OpenFWI subset: {counts}")


def _y_to_inches(value: float) -> float:
    return PLOT_TOP_IN + (Y_MAX - float(value)) / (Y_MAX - Y_MIN) * PLOT_HEIGHT_IN


def _rgb(hex_value: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_value)


def _add_text(
    slide: Any,
    text: str,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    size: float,
    color: str = TEXT_COLOR,
    bold: bool = False,
    italic: bool = False,
    align: str = "left",
    rotation: float = 0.0,
    name: str | None = None,
) -> Any:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if name:
        shape.name = name
    shape.rotation = rotation
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }[align]
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)
    return shape


def _add_line(
    slide: Any,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str,
    width: float = 0.7,
    dashed: bool = False,
    name: str | None = None,
) -> Any:
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches, Pt

    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    if name:
        line.name = name
    line.line.color.rgb = _rgb(color)
    line.line.width = Pt(width)
    if dashed:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return line


def _add_violin(
    slide: Any,
    values: np.ndarray,
    *,
    center_x: float,
    max_half_width: float,
    fill_color: str,
    line_color: str,
    name: str,
) -> Any:
    from pptx.util import Inches, Pt

    values = np.asarray(values, dtype=np.float64)
    low = max(Y_MIN, float(values.min()) - 0.025)
    high = min(Y_MAX, float(values.max()) + 0.025)
    grid = np.linspace(low, high, 150)
    if values.size < 2 or float(values.std()) < 1e-10:
        scale = max((high - low) / 6.0, 1e-3)
        density = np.exp(-0.5 * np.square((grid - float(values.mean())) / scale))
    else:
        density = gaussian_kde(values, bw_method="scott")(grid)
    density = density / max(float(density.max()), 1e-12)
    half_width = max_half_width * density
    left_points = [(center_x - float(width), _y_to_inches(value)) for width, value in zip(half_width, grid)]
    right_points = [
        (center_x + float(width), _y_to_inches(value))
        for width, value in zip(half_width[::-1], grid[::-1])
    ]
    points = left_points + right_points
    local_scale = float(Inches(1)) / 1000.0
    local_points = [(int(round(x * 1000)), int(round(y * 1000))) for x, y in points]
    builder = slide.shapes.build_freeform(
        start_x=local_points[0][0],
        start_y=local_points[0][1],
        scale=(local_scale, local_scale),
    )
    builder.add_line_segments(local_points[1:], close=True)
    shape = builder.convert_to_shape()
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill_color)
    shape.line.color.rgb = _rgb(line_color)
    shape.line.width = Pt(0.75)
    return shape


def _add_boxplot(slide: Any, values: np.ndarray, *, center_x: float, category_width: float, name: str) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    values = np.asarray(values, dtype=np.float64)
    q1, median, q3 = np.quantile(values, [0.25, 0.50, 0.75])
    iqr = q3 - q1
    lower_candidates = values[values >= q1 - 1.5 * iqr]
    upper_candidates = values[values <= q3 + 1.5 * iqr]
    lower = float(lower_candidates.min())
    upper = float(upper_candidates.max())
    box_half = category_width * 0.105
    cap_half = category_width * 0.10
    median_half = category_width * 0.21

    _add_line(slide, center_x, _y_to_inches(lower), center_x, _y_to_inches(upper), color=AXIS_COLOR, width=0.75, name=f"{name}_Whisker")
    _add_line(slide, center_x - cap_half, _y_to_inches(lower), center_x + cap_half, _y_to_inches(lower), color=AXIS_COLOR, width=0.75, name=f"{name}_LowerCap")
    _add_line(slide, center_x - cap_half, _y_to_inches(upper), center_x + cap_half, _y_to_inches(upper), color=AXIS_COLOR, width=0.75, name=f"{name}_UpperCap")

    top = _y_to_inches(q3)
    bottom = _y_to_inches(q1)
    box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(center_x - box_half),
        Inches(top),
        Inches(box_half * 2),
        Inches(bottom - top),
    )
    box.name = f"{name}_IQR"
    box.fill.solid()
    box.fill.fore_color.rgb = _rgb("FFFFFF")
    box.line.color.rgb = _rgb(AXIS_COLOR)
    box.line.width = Pt(0.8)
    _add_line(
        slide,
        center_x - median_half,
        _y_to_inches(median),
        center_x + median_half,
        _y_to_inches(median),
        color=TEXT_COLOR,
        width=1.25,
        name=f"{name}_Median",
    )


def export_transport_energy_pptx(
    *,
    qt_rows: Sequence[Mapping[str, Any]],
    metadata: Mapping[str, Any],
    output_path: str | Path,
    metadata_output_path: str | Path | None = None,
    expected_count: int = 33_600,
) -> Path:
    """Create a one-slide native-shape PowerPoint for the q_T diagnostic."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches

    _validate_rows(qt_rows, expected_count=expected_count)
    output_path = _resolve_path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    values_by_subset = {
        name: np.asarray([_qt_value(row) for row in qt_rows if str(row.get("dataset_name")) == name])
        for name in SUBSET_ORDER
    }
    summary = dict(metadata.get("q_t_summary", {}))
    required_summary = {
        "mean_q_t",
        "fraction_q_t_lt_1",
        "mean_q_t_ci_low",
        "mean_q_t_ci_high",
        "n",
    }
    if not required_summary.issubset(summary):
        raise ValueError(f"metadata q_t_summary is missing fields: {sorted(required_summary - set(summary))}")
    if int(summary["n"]) != len(qt_rows):
        raise ValueError(f"metadata reports n={summary['n']}, but CSV has {len(qt_rows)} rows")

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb("FFFFFF")

    _add_text(
        slide,
        "Full-Test Residual Transport-Energy Condition",
        left=0.42,
        top=0.10,
        width=12.45,
        height=0.42,
        size=21,
        bold=True,
        name="Title",
    )
    _add_text(
        slide,
        f"global q_T < 1: {float(summary['fraction_q_t_lt_1']):.1%}   |   "
        f"mean q_T: {float(summary['mean_q_t']):.3f} "
        f"[{float(summary['mean_q_t_ci_low']):.3f}, {float(summary['mean_q_t_ci_high']):.3f}]   |   "
        f"n = {int(summary['n']):,}",
        left=PLOT_LEFT_IN,
        top=0.57,
        width=9.7,
        height=0.30,
        size=13,
        color=MUTED_TEXT,
        name="GlobalSummary",
    )

    tick_values = np.arange(0.5, 1.11, 0.1)
    for tick in tick_values:
        y = _y_to_inches(float(tick))
        _add_line(slide, PLOT_LEFT_IN, y, PLOT_LEFT_IN + PLOT_WIDTH_IN, y, color=GRID_COLOR, width=0.55, name=f"Grid_{tick:.1f}")
        _add_text(
            slide,
            f"{tick:.1f}",
            left=0.54,
            top=y - 0.12,
            width=0.38,
            height=0.24,
            size=11.5,
            color=AXIS_COLOR,
            align="right",
            name=f"YTick_{tick:.1f}",
        )

    threshold_y = _y_to_inches(1.0)
    _add_line(
        slide,
        PLOT_LEFT_IN,
        threshold_y,
        PLOT_LEFT_IN + PLOT_WIDTH_IN,
        threshold_y,
        color=DIAGNOSTIC_PURPLE,
        width=1.1,
        dashed=True,
        name="Threshold_qT_1",
    )
    _add_line(slide, 11.50, 0.72, 11.92, 0.72, color=DIAGNOSTIC_PURPLE, width=1.1, dashed=True, name="ThresholdLegendLine")
    _add_text(slide, "q_T = 1", left=11.99, top=0.60, width=0.70, height=0.25, size=12.5, italic=True, name="ThresholdLegend")

    category_width = PLOT_WIDTH_IN / len(SUBSET_ORDER)
    for index, name in enumerate(SUBSET_ORDER):
        center_x = PLOT_LEFT_IN + category_width * (index + 0.5)
        values = values_by_subset[name]
        is_velocity = index < 4
        _add_violin(
            slide,
            values,
            center_x=center_x,
            max_half_width=category_width * 0.36,
            fill_color=BACKGROUND_FILL if is_velocity else STRUCTURE_FILL,
            line_color=BACKGROUND_LINE if is_velocity else STRUCTURE_LINE,
            name=f"Violin_{name}",
        )
        _add_boxplot(slide, values, center_x=center_x, category_width=category_width, name=f"Box_{name}")
        fraction = float(np.mean(values < 1.0))
        _add_text(
            slide,
            f"q_T < 1: {fraction:.0%}",
            left=center_x - category_width * 0.47,
            top=_y_to_inches(1.045) - 0.13,
            width=category_width * 0.94,
            height=0.25,
            size=10.5,
            color=MUTED_TEXT,
            align="center",
            name=f"Fraction_{name}",
        )
        _add_text(
            slide,
            name,
            left=center_x - category_width * 0.49,
            top=PLOT_TOP_IN + PLOT_HEIGHT_IN + 0.05,
            width=category_width * 0.98,
            height=0.30,
            size=11.5,
            align="center",
            name=f"Label_{name}",
        )

    _add_line(slide, PLOT_LEFT_IN, PLOT_TOP_IN, PLOT_LEFT_IN, PLOT_TOP_IN + PLOT_HEIGHT_IN, color=AXIS_COLOR, width=0.9, name="YAxis")
    _add_line(slide, PLOT_LEFT_IN, PLOT_TOP_IN + PLOT_HEIGHT_IN, PLOT_LEFT_IN + PLOT_WIDTH_IN, PLOT_TOP_IN + PLOT_HEIGHT_IN, color=AXIS_COLOR, width=0.9, name="XAxis")
    _add_text(
        slide,
        "Normalized transport burden q_T",
        left=0.02,
        top=1.55,
        width=0.36,
        height=2.30,
        size=13,
        rotation=270,
        align="center",
        name="YAxisTitle",
    )
    _add_text(
        slide,
        "OpenFWI subset",
        left=5.35,
        top=4.72,
        width=2.65,
        height=0.31,
        size=13.5,
        align="center",
        name="XAxisTitle",
    )

    presentation.save(output_path)
    reopened = Presentation(output_path)
    if len(reopened.slides) != 1:
        raise RuntimeError(f"Editable figure deck must contain one slide, found {len(reopened.slides)}")
    shapes = list(reopened.slides[0].shapes)
    picture_count = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes)
    freeform_count = sum(shape.shape_type == MSO_SHAPE_TYPE.FREEFORM for shape in shapes)
    if picture_count != 0 or freeform_count != len(SUBSET_ORDER):
        raise RuntimeError(
            f"PPTX editability check failed: pictures={picture_count}, freeforms={freeform_count}"
        )

    if metadata_output_path is not None:
        metadata_output = _resolve_path(metadata_output_path)
        metadata_output.parent.mkdir(parents=True, exist_ok=True)
        metadata_output.write_text(
            json.dumps(
                {
                    "presentation": str(output_path),
                    "source_qt_records": len(qt_rows),
                    "subset_order": list(SUBSET_ORDER),
                    "slide_size_inches": [SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN],
                    "shape_count": len(shapes),
                    "freeform_violin_count": freeform_count,
                    "picture_count": picture_count,
                    "all_elements_native_editable": picture_count == 0,
                    "font": FONT_NAME,
                },
                indent=2,
                sort_keys=True,
            ),
            encoding="utf-8",
        )
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Export the q_T diagnostic as editable native PowerPoint shapes.")
    parser.add_argument("--qt-samples", default=str(DEFAULT_OUTPUT_DIR / "qt_samples.csv"))
    parser.add_argument("--metadata", default=str(DEFAULT_OUTPUT_DIR / "metadata.json"))
    parser.add_argument(
        "--output",
        default=str(DEFAULT_OUTPUT_DIR / "figure4_transport_energy_editable.pptx"),
    )
    parser.add_argument(
        "--output-metadata",
        default=str(DEFAULT_OUTPUT_DIR / "figure4_transport_energy_pptx_metadata.json"),
    )
    args = parser.parse_args()
    output = export_transport_energy_pptx(
        qt_rows=_read_csv(args.qt_samples),
        metadata=_read_json(args.metadata),
        output_path=args.output,
        metadata_output_path=args.output_metadata,
    )
    print(json.dumps({"editable_pptx": str(output)}, indent=2))


if __name__ == "__main__":
    main()
