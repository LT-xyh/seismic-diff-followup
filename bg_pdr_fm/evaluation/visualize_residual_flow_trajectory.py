"""Render a real PD-BG-RFM residual Flow Matching trajectory.

This script is checkpoint-first. It can select either a record from the
official held-out evaluation outputs or an explicitly requested all-split
visualization asset, runs the formal checkpoint with the production
conditioning path, and records the exact 50-step Euler states, vector fields,
and increments.

The script does not use the legacy post-hoc RFW visualizer and does not create
interpolated or warped proxy states.
"""

from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import io
import json
import os
import random
import sys
from pathlib import Path
from typing import Any, Iterable

import lightning
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import torch
from matplotlib.colors import Normalize, TwoSlopeNorm
from matplotlib.cm import ScalarMappable
from omegaconf import OmegaConf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from bg_pdr_fm.data import batch_to_device, collate_bg_samples, validate_bg_batch
from bg_pdr_fm.lightning import BGPDRFMLightning
from bg_pdr_fm.lightning.stage_losses import select_residual_background
from bg_pdr_fm.runtime import configure_torch_runtime
from bg_pdr_fm.training.train_bg_pdr_fm import build_dataset


DEFAULT_CONFIG = (
    REPO_ROOT
    / "bg_pdr_fm/configs/openfwi_lmdb_joint_full_contrastive_bgfm_pixelfm_predbg_e100.yaml"
)
DEFAULT_CHECKPOINT = (
    REPO_ROOT
    / "logs/bg_pdr_fm/joint_full_contrastive_bgfm_pixelfm_predbg_e100/lightning/checkpoints/"
    "joint_full_contrastive_bgfm_pixelfm_predbg_e100-epoch_99-loss0.1036.ckpt"
)
DEFAULT_EVAL_DIR = (
    REPO_ROOT
    / "logs/bg_pdr_fm/joint_full_contrastive_bgfm_pixelfm_predbg_e100/"
    "evaluation_epoch99_global_heldout_full_parallel_bs64"
)
DEFAULT_OUTPUT_DIR = REPO_ROOT / "docs/paper/AAAI2027/figures/residual_flow_trajectory"
TARGET_DATASET = "FlatFaultB"
TARGET_SOURCE_SAMPLE_INDEX = 27006
TRACE_SEED = 2027
TRACE_STEPS = 50
STATE_INDICES = (0, 10, 25, 40, 50)
VELOCITY_INDICES = (0, 10, 25, 40, 49)
STATE_LABELS = (r"$R_0=\xi$", r"$R_{10}$", r"$R_{25}$", r"$R_{40}$", r"$R_{50}=\hat R$")
VELOCITY_LABELS = (r"$v_0$", r"$v_{10}$", r"$v_{25}$", r"$v_{40}$", r"$v_{49}$")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _identity(row: dict[str, str]) -> tuple[int, str, int]:
    return int(row["dataset_id"]), str(row["dataset_name"]), int(row["source_sample_index"])


def select_representative_record(
    manifest_path: Path,
    metrics_path: Path,
    dataset_name: str = TARGET_DATASET,
) -> tuple[dict[str, str], dict[str, str], float]:
    """Select the official record closest to the subset SSIM median."""
    manifest = _read_csv(manifest_path)
    metrics = _read_csv(metrics_path)
    if not manifest or not metrics:
        raise ValueError("The official held-out manifest and merged metrics must both be non-empty.")
    manifest_ids = {_identity(row) for row in manifest}
    metric_ids = {_identity(row) for row in metrics}
    if manifest_ids != metric_ids:
        raise ValueError(
            "Official manifest and metrics identity sets differ: "
            f"manifest={len(manifest_ids)}, metrics={len(metric_ids)}."
        )
    subset = [row for row in metrics if row["dataset_name"] == dataset_name]
    if not subset:
        raise ValueError(f"No official metrics found for dataset {dataset_name!r}.")
    median_ssim = float(np.median([float(row["ssim"]) for row in subset]))
    selected_metrics = min(
        subset,
        key=lambda row: (
            abs(float(row["ssim"]) - median_ssim),
            int(row["source_sample_index"]),
        ),
    )
    selected_manifest = next(
        row for row in manifest if _identity(row) == _identity(selected_metrics)
    )
    return selected_manifest, selected_metrics, median_ssim


def select_requested_record(
    manifest_path: Path,
    metrics_path: Path,
    *,
    dataset_name: str,
    source_sample_index: int | None,
    dataset_id: int | None = None,
) -> dict[str, Any]:
    """Resolve either the official median record or an explicit asset record.

    Explicit records outside the official held-out manifest are intentionally
    marked as such and receive no fabricated evaluation metrics.
    """
    if source_sample_index is None:
        selected_manifest, selected_metrics, median_ssim = select_representative_record(
            manifest_path,
            metrics_path,
            dataset_name=dataset_name,
        )
        return {
            "selected_manifest": selected_manifest,
            "selected_metrics": selected_metrics,
            "median_ssim": median_ssim,
            "selection_rule": f"{dataset_name} official PD-BG-RFM SSIM closest to subset median; ties by source_sample_index",
            "provenance_scope": "official_global_heldout",
            "dataset_split": "test",
        }

    if dataset_id is None:
        raise ValueError("dataset_id is required when selecting an explicit source sample index.")
    manifest = _read_csv(manifest_path)
    metrics = _read_csv(metrics_path)
    heldout_manifest = [
        row
        for row in manifest
        if row["dataset_name"] == dataset_name
        and int(row["source_sample_index"]) == int(source_sample_index)
    ]
    heldout_metrics = [
        row
        for row in metrics
        if row["dataset_name"] == dataset_name
        and int(row["source_sample_index"]) == int(source_sample_index)
    ]
    if heldout_manifest:
        if len(heldout_manifest) != 1 or len(heldout_metrics) != 1:
            raise ValueError(
                f"Expected one official record for {dataset_name}/{source_sample_index}, "
                f"found manifest={len(heldout_manifest)}, metrics={len(heldout_metrics)}."
            )
        subset = [row for row in metrics if row["dataset_name"] == dataset_name]
        median_ssim = float(np.median([float(row["ssim"]) for row in subset]))
        return {
            "selected_manifest": heldout_manifest[0],
            "selected_metrics": heldout_metrics[0],
            "median_ssim": median_ssim,
            "selection_rule": f"user-requested {dataset_name} source_sample_index={source_sample_index}; record is in official global held-out manifest",
            "provenance_scope": "official_global_heldout",
            "dataset_split": "test",
        }

    selected_manifest = {
        "dataset_id": str(dataset_id),
        "dataset_name": dataset_name,
        "source_sample_index": str(int(source_sample_index)),
        "worker": "user_selected_figure2_asset",
    }
    selected_metrics = {
        "sample_index": "",
        "batch_index": "",
        "item_index": "",
        "mae": "",
        "rmse": "",
        "ssim": "",
        "dataset_id": str(dataset_id),
        "dataset_name": dataset_name,
        "source_sample_index": str(int(source_sample_index)),
        "provenance_note": "not present in official global held-out manifest; no official metrics reported",
    }
    return {
        "selected_manifest": selected_manifest,
        "selected_metrics": selected_metrics,
        "median_ssim": None,
        "selection_rule": f"user-specified Fig. 2 asset: {dataset_name} source_sample_index={source_sample_index}; not in official global held-out manifest",
        "provenance_scope": "not_in_official_global_heldout",
        "dataset_split": "all",
    }


def _set_seed(seed: int) -> None:
    random.seed(seed)
    np.random.seed(seed)
    lightning.seed_everything(int(seed), workers=True)
    torch.manual_seed(int(seed))
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(int(seed))


def _prepare_config(config_path: Path):
    conf = OmegaConf.load(config_path)
    conf = copy.deepcopy(conf)
    conf.training.stage = "joint_full"
    conf.training.load_stage_checkpoint = None
    if "joint" in conf.training and "warm_start_checkpoints" in conf.training.joint:
        for label in ("contrastive", "background", "residual"):
            conf.training.joint.warm_start_checkpoints[label] = None
    conf.data.well_random = False
    return conf


def _load_full_checkpoint(model: BGPDRFMLightning, checkpoint_path: Path) -> dict[str, Any]:
    state = torch.load(checkpoint_path, map_location="cpu", weights_only=False)
    if not isinstance(state, dict) or "state_dict" not in state:
        raise KeyError(f"Checkpoint {checkpoint_path} does not contain a state_dict.")
    incompatible = model.load_state_dict(state["state_dict"], strict=False)
    if incompatible.missing_keys or incompatible.unexpected_keys:
        raise RuntimeError(
            "Formal checkpoint load was not exact: "
            f"missing={incompatible.missing_keys}, unexpected={incompatible.unexpected_keys}"
        )
    return {
        "path": str(checkpoint_path),
        "sha256": sha256_file(checkpoint_path),
        "epoch": state.get("epoch"),
        "global_step": state.get("global_step"),
        "missing_keys": list(incompatible.missing_keys),
        "unexpected_keys": list(incompatible.unexpected_keys),
        "state_dict_keys": len(state["state_dict"]),
    }


def _find_dataset_record(dataset: Any, selected_manifest: dict[str, str]) -> tuple[int, dict[str, Any]]:
    wanted = (
        int(selected_manifest["dataset_id"]),
        str(selected_manifest["dataset_name"]),
        int(selected_manifest["source_sample_index"]),
    )
    matches = [
        (index, record)
        for index, record in enumerate(dataset.records)
        if (
            int(record["dataset_id"]),
            str(record["dataset_name"]),
            int(record["sample_index"]),
        )
        == wanted
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one dataset record for {wanted}, found {len(matches)}.")
    return matches[0]


def _infer_dataset_id(dataset: Any, dataset_name: str, source_sample_index: int) -> int:
    """Resolve the stable subset id from the all-split dataset index."""
    matches = [
        record
        for record in dataset.records
        if str(record["dataset_name"]) == str(dataset_name)
        and int(record["sample_index"]) == int(source_sample_index)
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one all-split record for {dataset_name}/{source_sample_index}, "
            f"found {len(matches)}."
        )
    return int(matches[0]["dataset_id"])


@torch.no_grad()
def run_real_trace(
    conf: Any,
    checkpoint_path: Path,
    dataset: Any,
    selected_manifest: dict[str, str],
    *,
    device: torch.device,
    seed: int,
    steps: int,
) -> dict[str, Any]:
    dataset_index, _ = _find_dataset_record(dataset, selected_manifest)
    item = dataset[dataset_index]
    batch = collate_bg_samples([item])
    validate_bg_batch(batch, context="residual flow trajectory selected sample")
    batch = batch_to_device(batch, device)

    model = BGPDRFMLightning(conf)
    checkpoint_info = _load_full_checkpoint(model, checkpoint_path)
    model.eval().to(device=device, dtype=torch.float32)

    features, conditions, background, _ = model._common_forward(batch, use_background_target=False)
    del features
    if conditions is None or background is None or model.residual is None or model.codec is None:
        raise RuntimeError("The formal joint_full checkpoint did not expose the generation stack.")
    if str(model.codec_type).lower() not in {"identity", "pixel", "none"}:
        raise RuntimeError(
            "This visualization is defined for the formal identity/pixel-space residual model; "
            f"got codec_type={model.codec_type!r}."
        )

    residual_background = select_residual_background(
        model.conf,
        model.filter,
        batch.depth_vel,
        background.bg_hat.detach(),
    )
    z_bg = model.codec.encode(residual_background.detach(), deterministic=True)
    cond_structural = conditions.structural
    background_context = model.residual.background_context(
        residual_background.detach(),
        out_hw=tuple(z_bg.shape[-2:]),
        z_bg=z_bg,
    )
    if model.residual.condition_merge != "concat":
        raise RuntimeError(
            "The formal residual condition must be concatenated for this figure, "
            f"got condition_merge={model.residual.condition_merge!r}."
        )
    merged_condition = torch.cat([cond_structural, background_context], dim=1)

    _set_seed(seed)
    noise = torch.randn_like(z_bg)
    trace = model.residual.sample_with_trace(
        cond_structural,
        residual_background.detach(),
        x_size=tuple(z_bg.shape),
        steps=steps,
        z_bg=z_bg,
        noise=noise,
    )
    production_sample = model.residual.sample(
        cond_structural,
        residual_background.detach(),
        x_size=tuple(z_bg.shape),
        steps=steps,
        z_bg=z_bg,
        noise=noise,
    )
    torch.testing.assert_close(trace["states"][-1], production_sample)
    torch.testing.assert_close(trace["condition"], merged_condition)

    residual_hat = trace["states"][-1]
    velocity_hat = model.codec.decode(z_bg + residual_hat, out_hw=tuple(batch.depth_vel.shape[-2:]))
    composed_velocity = residual_background + residual_hat
    torch.testing.assert_close(velocity_hat, composed_velocity)

    states = trace["states"][:, 0].detach().float().cpu().numpy()
    velocities = trace["velocities"][:, 0].detach().float().cpu().numpy()
    deltas = trace["deltas"][:, 0].detach().float().cpu().numpy()
    times = trace["times"][:, 0].detach().float().cpu().numpy()
    arrays = {
        "residual_states": states,
        "velocity_fields": velocities,
        "euler_deltas": deltas,
        "times": times,
        "noise": noise[0].detach().float().cpu().numpy(),
        "condition_structural": cond_structural[0].detach().float().cpu().numpy(),
        "background_context": background_context[0].detach().float().cpu().numpy(),
        "merged_condition": merged_condition[0].detach().float().cpu().numpy(),
        "channel_rms_structural": cond_structural[0].square().mean(dim=0).sqrt().detach().float().cpu().numpy(),
        "channel_rms_background_context": background_context[0].square().mean(dim=0).sqrt().detach().float().cpu().numpy(),
        "background_hat": residual_background[0].detach().float().cpu().numpy(),
        "residual_hat": residual_hat[0].detach().float().cpu().numpy(),
        "velocity_hat": velocity_hat[0].detach().float().cpu().numpy(),
        "target_velocity": batch.depth_vel[0].detach().float().cpu().numpy(),
    }
    for name, array in arrays.items():
        if not np.isfinite(array).all():
            raise FloatingPointError(f"Non-finite values found in {name}.")

    if states.shape[1:] != (1, 70, 70):
        raise ValueError(f"Expected residual state shape [K, 1, 70, 70], got {states.shape}.")
    if velocities.shape[1:] != (1, 70, 70) or deltas.shape[1:] != (1, 70, 70):
        raise ValueError(
            "Expected velocity and delta shapes [K, 1, 70, 70], "
            f"got velocity={velocities.shape}, delta={deltas.shape}."
        )
    max_trace_error = float(np.max(np.abs(states[-1] - production_sample[0].detach().float().cpu().numpy())))
    max_composition_error = float(
        np.max(np.abs(arrays["velocity_hat"] - (arrays["background_hat"] + arrays["residual_hat"])))
    )
    return {
        "arrays": arrays,
        "checkpoint_info": checkpoint_info,
        "dataset_index": dataset_index,
        "selected_item": {
            key: value.detach().cpu().tolist()
            for key, value in item.items()
            if key in {"dataset_id", "sample_index", "modality_mask", "modality_quality"}
        },
        "validation": {
            "trace_final_vs_production_max_abs": max_trace_error,
            "composition_max_abs": max_composition_error,
            "trace_final_vs_production_assert_close": True,
            "composition_assert_close": True,
            "all_arrays_finite": True,
            "state_count": int(states.shape[0]),
            "velocity_count": int(velocities.shape[0]),
            "delta_count": int(deltas.shape[0]),
        },
    }


def _symmetric_limit(values: Iterable[np.ndarray], percentile: float = 100.0) -> float:
    flat = np.concatenate([np.asarray(value).reshape(-1) for value in values])
    limit = float(np.percentile(np.abs(flat), percentile))
    return max(limit, 1e-8)


def _display_image(image: np.ndarray) -> np.ndarray:
    image = np.asarray(image)
    if image.ndim == 3 and image.shape[0] == 1:
        return image[0]
    if image.ndim == 3 and image.shape[-1] == 1:
        return image[..., 0]
    if image.ndim != 2:
        raise ValueError(f"Expected a single-channel 2D image for display, got shape {image.shape}.")
    return image


def _add_image_axis(axis, image: np.ndarray, cmap: str, norm: Normalize, label: str, *, fontsize: float = 6.2):
    axis.imshow(_display_image(image), cmap=cmap, norm=norm, origin="upper", interpolation="nearest", aspect="equal")
    axis.set_xticks([])
    axis.set_yticks([])
    axis.set_title(label, fontsize=fontsize, pad=2.0)
    for spine in axis.spines.values():
        spine.set_linewidth(0.45)
        spine.set_color("#444444")
    return axis


def _add_horizontal_colorbar(fig, rect, cmap: str, norm: Normalize, label: str):
    axis = fig.add_axes(rect)
    sm = ScalarMappable(norm=norm, cmap=cmap)
    colorbar = fig.colorbar(sm, cax=axis, orientation="horizontal")
    colorbar.ax.tick_params(labelsize=5, length=2, pad=1)
    colorbar.set_label(label, fontsize=5.5, labelpad=1)
    return colorbar


def render_figure(arrays: dict[str, np.ndarray], output_dir: Path) -> dict[str, Any]:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 6.5,
            "mathtext.fontset": "stix",
            "axes.unicode_minus": False,
            "savefig.facecolor": "white",
            "figure.facecolor": "white",
        }
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    states = arrays["residual_states"]
    velocities = arrays["velocity_fields"]
    state_limit = _symmetric_limit([states])
    velocity_limit = _symmetric_limit([velocities])
    feature_limit = float(
        np.percentile(
            np.concatenate(
                [
                    arrays["channel_rms_structural"].reshape(-1),
                    arrays["channel_rms_background_context"].reshape(-1),
                ]
            ),
            99.5,
        )
    )
    feature_limit = max(feature_limit, 1e-8)
    speed_limit = _symmetric_limit([arrays["background_hat"], arrays["velocity_hat"]])
    residual_limit = _symmetric_limit([arrays["residual_hat"]])

    state_norm = TwoSlopeNorm(vmin=-state_limit, vcenter=0.0, vmax=state_limit)
    velocity_norm = TwoSlopeNorm(vmin=-velocity_limit, vcenter=0.0, vmax=velocity_limit)
    residual_norm = TwoSlopeNorm(vmin=-residual_limit, vcenter=0.0, vmax=residual_limit)
    speed_norm = Normalize(vmin=-speed_limit, vmax=speed_limit)
    feature_norm = Normalize(vmin=0.0, vmax=feature_limit)

    fig = plt.figure(figsize=(7.35, 4.75))
    outer = fig.add_gridspec(
        1,
        3,
        width_ratios=(1.28, 4.55, 1.32),
        left=0.035,
        right=0.985,
        bottom=0.105,
        top=0.895,
        wspace=0.24,
    )
    fig.suptitle(
        "Prediction-Consistent Residual Flow Trajectory",
        fontsize=10.0,
        fontweight="bold",
        y=0.965,
    )

    panel_a = outer[0].subgridspec(4, 1, height_ratios=(0.28, 1.0, 1.0, 0.56), hspace=0.28)
    ax = fig.add_subplot(panel_a[0])
    ax.axis("off")
    ax.text(0.0, 0.5, "(a) Conditioning", fontsize=8.0, fontweight="bold", va="center")
    ax = fig.add_subplot(panel_a[1])
    _add_image_axis(ax, arrays["channel_rms_structural"], "viridis", feature_norm, r"$C_{\mathrm{str}}$  (channel RMS)")
    ax = fig.add_subplot(panel_a[2])
    _add_image_axis(
        ax,
        arrays["channel_rms_background_context"],
        "cividis",
        feature_norm,
        r"$\psi_{\mathrm{bg}}(\hat B)$  (channel RMS)",
    )
    ax = fig.add_subplot(panel_a[3])
    ax.axis("off")
    ax.text(
        0.5,
        0.68,
        r"$c_R=[C_{\mathrm{str}}\,;\,\psi_{\mathrm{bg}}(\hat B)]$",
        ha="center",
        va="center",
        fontsize=7.0,
        bbox={"boxstyle": "round,pad=0.32", "fc": "#f4f7f8", "ec": "#536878", "lw": 0.65},
    )
    ax.text(0.5, 0.12, "feature energy; no velocity units", ha="center", va="center", fontsize=5.0, color="#4a4a4a")

    panel_b = outer[1].subgridspec(3, 5, height_ratios=(1.0, 1.0, 0.22), hspace=0.33, wspace=0.16)
    state_axes = []
    field_axes = []
    for col, (state_idx, label) in enumerate(zip(STATE_INDICES, STATE_LABELS)):
        axis = fig.add_subplot(panel_b[0, col])
        state_axes.append(axis)
        _add_image_axis(axis, states[state_idx, 0], "RdBu_r", state_norm, label)
        if col < len(STATE_INDICES) - 1:
            axis.text(1.065, 0.50, r"$\rightarrow$", transform=axis.transAxes, ha="center", va="center", fontsize=10.0)
    for col, (velocity_idx, label) in enumerate(zip(VELOCITY_INDICES, VELOCITY_LABELS)):
        axis = fig.add_subplot(panel_b[1, col])
        field_axes.append(axis)
        _add_image_axis(axis, velocities[velocity_idx, 0], "RdBu_r", velocity_norm, label)
    field_axes[0].set_ylabel(r"$v_\theta(R_k,t_k,c_R)$", fontsize=6.2, labelpad=4)
    state_axes[0].set_ylabel("state", fontsize=6.2, labelpad=4)
    title_axis = fig.add_subplot(panel_b[2, :])
    title_axis.axis("off")
    title_axis.text(
        0.5,
        0.50,
        r"$R_{k+1}=R_k+\Delta t\,v_k,\quad \Delta t=1/50$"
        "\n"
        r"$R_t=(1-t)\xi+t(V-\hat B)$  (training target; held-out $V$ is not an inference input)",
        ha="center",
        va="center",
        fontsize=6.0,
        color="#333333",
    )

    panel_c = outer[2].subgridspec(4, 1, height_ratios=(0.28, 1.0, 1.0, 1.0), hspace=0.28)
    ax = fig.add_subplot(panel_c[0])
    ax.axis("off")
    ax.text(0.0, 0.5, "(c) Composition", fontsize=8.0, fontweight="bold", va="center")
    ax = fig.add_subplot(panel_c[1])
    _add_image_axis(ax, arrays["background_hat"], "viridis", speed_norm, r"$\hat B$")
    ax = fig.add_subplot(panel_c[2])
    _add_image_axis(ax, arrays["residual_hat"], "RdBu_r", residual_norm, r"$\hat R=R_{50}$")
    ax = fig.add_subplot(panel_c[3])
    _add_image_axis(ax, arrays["velocity_hat"], "viridis", speed_norm, r"$\hat V=\hat B+\hat R$")

    _add_horizontal_colorbar(fig, (0.055, 0.055, 0.18, 0.016), "viridis", feature_norm, "feature RMS")
    _add_horizontal_colorbar(fig, (0.405, 0.055, 0.22, 0.016), "RdBu_r", state_norm, "residual state")
    _add_horizontal_colorbar(fig, (0.655, 0.055, 0.22, 0.016), "RdBu_r", velocity_norm, "predicted field")

    outputs = {}
    for suffix, kwargs in (
        ("pdf", {"bbox_inches": "tight", "pad_inches": 0.03}),
        ("svg", {"bbox_inches": "tight", "pad_inches": 0.03}),
        ("png", {"dpi": 360, "bbox_inches": "tight", "pad_inches": 0.03}),
    ):
        path = output_dir / f"residual_flow_trajectory.{suffix}"
        fig.savefig(path, **kwargs)
        outputs[suffix] = str(path)
    plt.close(fig)
    return {
        "files": outputs,
        "color_limits": {
            "feature_rms": [0.0, feature_limit],
            "residual_state": [-state_limit, state_limit],
            "velocity_field": [-velocity_limit, velocity_limit],
            "background_and_final_velocity": [-speed_limit, speed_limit],
            "final_residual": [-residual_limit, residual_limit],
        },
    }


def _ppt_text(slide, left, top, width, height, text, *, size=8.0, bold=False, color=(35, 35, 35), align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Pt(0)
    frame.margin_right = Pt(0)
    frame.margin_top = Pt(0)
    frame.margin_bottom = Pt(0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return shape


def _ppt_picture_bytes(image: np.ndarray, cmap: str, norm: Normalize) -> io.BytesIO:
    fig, axis = plt.subplots(figsize=(1.0, 1.0), dpi=220)
    axis.imshow(_display_image(image), cmap=cmap, norm=norm, origin="upper", interpolation="nearest")
    axis.set_axis_off()
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png", dpi=220, bbox_inches="tight", pad_inches=0)
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _ppt_picture(slide, image: np.ndarray, cmap: str, norm: Normalize, left, top, width, height, label: str):
    slide.shapes.add_picture(
        _ppt_picture_bytes(image, cmap, norm),
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    _ppt_text(slide, left, top + height + 0.015, width, 0.18, label, size=6.1)


def export_editable_pptx(arrays: dict[str, np.ndarray], output_path: Path) -> None:
    states = arrays["residual_states"]
    velocities = arrays["velocity_fields"]
    state_limit = _symmetric_limit([states])
    velocity_limit = _symmetric_limit([velocities])
    feature_limit = max(
        float(
            np.percentile(
                np.concatenate(
                    [
                        arrays["channel_rms_structural"].reshape(-1),
                        arrays["channel_rms_background_context"].reshape(-1),
                    ]
                ),
                99.5,
            )
        ),
        1e-8,
    )
    speed_limit = _symmetric_limit([arrays["background_hat"], arrays["velocity_hat"]])
    residual_limit = _symmetric_limit([arrays["residual_hat"]])
    state_norm = TwoSlopeNorm(vmin=-state_limit, vcenter=0.0, vmax=state_limit)
    velocity_norm = TwoSlopeNorm(vmin=-velocity_limit, vcenter=0.0, vmax=velocity_limit)
    residual_norm = TwoSlopeNorm(vmin=-residual_limit, vcenter=0.0, vmax=residual_limit)
    speed_norm = Normalize(vmin=-speed_limit, vmax=speed_limit)
    feature_norm = Normalize(vmin=0.0, vmax=feature_limit)

    prs = Presentation()
    prs.slide_width = Inches(7.35)
    prs.slide_height = Inches(4.75)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _ppt_text(slide, 0.1, 0.04, 7.1, 0.25, "Prediction-Consistent Residual Flow Trajectory", size=12.0, bold=True)
    _ppt_text(slide, 0.1, 0.33, 1.1, 0.2, "(a) Conditioning", size=8.3, bold=True, align=PP_ALIGN.LEFT)
    _ppt_picture(slide, arrays["channel_rms_structural"], "viridis", feature_norm, 0.10, 0.62, 0.88, 0.88, "C_str (channel RMS)")
    _ppt_picture(slide, arrays["channel_rms_background_context"], "cividis", feature_norm, 0.10, 1.72, 0.88, 0.88, "psi_bg(B_hat)")
    _ppt_text(slide, 0.05, 2.82, 1.0, 0.35, "c_R = [C_str ; psi_bg(B_hat)]", size=6.3, bold=True)

    _ppt_text(slide, 1.22, 0.33, 4.55, 0.2, "(b) Euler trajectory", size=8.3, bold=True, align=PP_ALIGN.LEFT)
    _ppt_text(slide, 1.22, 0.53, 4.55, 0.20, "R_(k+1) = R_k + dt v_k,   dt = 1/50", size=6.3, color=(50, 50, 50), align=PP_ALIGN.LEFT)
    x0, y0, tile, gap = 1.25, 0.86, 0.73, 0.16
    for col, (idx, label) in enumerate(zip(STATE_INDICES, STATE_LABELS)):
        x = x0 + col * (tile + gap)
        _ppt_picture(slide, states[idx, 0], "RdBu_r", state_norm, x, y0, tile, tile, label.replace("$", ""))
        if col < 4:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(x + tile + 0.025),
                Inches(y0 + 0.30),
                Inches(0.12),
                Inches(0.12),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(80, 80, 80)
            arrow.line.fill.background()
    for col, (idx, label) in enumerate(zip(VELOCITY_INDICES, VELOCITY_LABELS)):
        x = x0 + col * (tile + gap)
        _ppt_picture(slide, velocities[idx, 0], "RdBu_r", velocity_norm, x, 2.05, tile, tile, label.replace("$", ""))
    _ppt_text(slide, 1.22, 3.00, 4.55, 0.34, "R_t = (1-t) xi + t(V - B_hat);  v_theta(R_t,t,c_R)", size=6.3, color=(50, 50, 50), align=PP_ALIGN.LEFT)

    _ppt_text(slide, 5.95, 0.33, 1.2, 0.2, "(c) Composition", size=8.3, bold=True, align=PP_ALIGN.LEFT)
    _ppt_picture(slide, arrays["background_hat"], "viridis", speed_norm, 6.02, 0.70, 0.90, 0.90, "B_hat")
    _ppt_picture(slide, arrays["residual_hat"], "RdBu_r", residual_norm, 6.02, 1.88, 0.90, 0.90, "R_hat")
    _ppt_picture(slide, arrays["velocity_hat"], "viridis", speed_norm, 6.02, 3.06, 0.90, 0.90, "V_hat = B_hat + R_hat")
    _ppt_text(
        slide,
        0.12,
        4.45,
        7.05,
        0.18,
        "Feature maps are channel RMS (no velocity units); tile images are embedded data assets, while labels, equations, and arrows remain editable.",
        size=5.2,
        color=(80, 80, 80),
        align=PP_ALIGN.LEFT,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def write_outputs(
    output_dir: Path,
    *,
    arrays: dict[str, np.ndarray],
    trace_result: dict[str, Any],
    selected_manifest: dict[str, str],
    selected_metrics: dict[str, str],
    median_ssim: float | None,
    selection_info: dict[str, Any],
    config_path: Path,
    checkpoint_path: Path,
    eval_dir: Path,
    seed: int,
    steps: int,
    figure_info: dict[str, Any],
    visible_devices: str | None,
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    np.savez_compressed(output_dir / "trajectory_arrays.npz", **arrays)
    provenance_scope = str(selection_info["provenance_scope"])
    dataset_split = str(selection_info["dataset_split"])
    selection_rule = str(selection_info["selection_rule"])
    has_official_metrics = provenance_scope == "official_global_heldout"
    sample_row = {
        **selected_manifest,
        "dataset_split": dataset_split,
        "provenance_scope": provenance_scope,
        "official_ssim": selected_metrics.get("ssim", "") if has_official_metrics else "",
        "official_rmse": selected_metrics.get("rmse", "") if has_official_metrics else "",
        "official_mae": selected_metrics.get("mae", "") if has_official_metrics else "",
        "subset_ssim_median": f"{median_ssim:.12f}" if median_ssim is not None else "",
        "absolute_distance_to_median": (
            f"{abs(float(selected_metrics['ssim']) - median_ssim):.12f}"
            if has_official_metrics and median_ssim is not None and selected_metrics.get("ssim", "")
            else ""
        ),
        "selection_rule": selection_rule,
    }
    with (output_dir / "sample_selection.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(sample_row))
        writer.writeheader()
        writer.writerow(sample_row)

    metadata = {
        "figure_title": "Prediction-Consistent Residual Flow Trajectory",
        "config_path": str(config_path),
        "config_sha256": sha256_file(config_path),
        "checkpoint": trace_result["checkpoint_info"],
        "official_evaluation_dir": str(eval_dir),
        "official_manifest": str(eval_dir / "held_out_manifest.csv"),
        "official_metrics": str(eval_dir / "merged_metrics.csv"),
        "selected_manifest_record": selected_manifest,
        "selected_official_metrics": selected_metrics if has_official_metrics else None,
        "selected_evaluation_metrics_available": has_official_metrics,
        "provenance_scope": provenance_scope,
        "dataset_split": dataset_split,
        "selection_rule": selection_rule,
        "subset_ssim_median": median_ssim,
        "dataset_index": trace_result["dataset_index"],
        "trace_seed": seed,
        "steps": steps,
        "delta_t": 1.0 / float(steps),
        "cuda_visible_devices": visible_devices,
        "model": {
            "codec_type": "identity",
            "residual_backend": "unet_fm_film",
            "residual_condition_merge": "concat",
            "residual_background_context_source": "codec_latent",
            "latent_hw": [70, 70],
            "condition_formula": "c_R = [C_str; psi_bg(B_hat)]",
            "path_formula": "R_t = (1-t) xi + t(V-B_hat)",
            "euler_formula": "R_{k+1} = R_k + (1/50) v_theta(R_k,t_k,c_R)",
        },
        "array_shapes": {key: list(value.shape) for key, value in arrays.items()},
        "figure_files": figure_info["files"],
        "color_limits": figure_info["color_limits"],
        "validation": trace_result["validation"],
    }
    (output_dir / "trajectory_metadata.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")

    if has_official_metrics:
        metric_lines = f"""- Official SSIM: {selected_metrics.get("ssim", "")}
- Subset SSIM median: {median_ssim:.12f}
- Official evaluation row: sample_index={selected_metrics.get("sample_index", "")}, batch_index={selected_metrics.get("batch_index", "")}, item_index={selected_metrics.get("item_index", "")}"""
    else:
        metric_lines = """- Official metrics: not reported; this record is outside the official global held-out manifest.
- Evaluation row: none (visualization asset only)."""

    readme = f"""# Prediction-Consistent Residual Flow Trajectory

This directory contains a real trajectory captured from the formal PD-BG-RFM
checkpoint. No proxy trajectory, interpolation-only state, spatial warp, or
post-hoc RFW registration was used.

## Provenance

- Checkpoint: {checkpoint_path}
- Checkpoint SHA256: {trace_result["checkpoint_info"]["sha256"]}
- Config: {config_path}
- Official held-out manifest: {eval_dir / "held_out_manifest.csv"}
- Official metrics: {eval_dir / "merged_metrics.csv"}
- Selected subset: {selected_manifest["dataset_name"]}
- Dataset split used to load the asset: {dataset_split}
- Provenance scope: {provenance_scope}
- Selection rule: {selection_rule}
- Selected dataset index: {trace_result["dataset_index"]}
- Selected source sample index: {selected_manifest["source_sample_index"]}
- Record ID: dataset_id={selected_manifest["dataset_id"]}, dataset_name={selected_manifest["dataset_name"]}, source_sample_index={selected_manifest["source_sample_index"]}
{metric_lines}
- Trace seed: {seed}
- Euler steps: {steps}
- CUDA_VISIBLE_DEVICES: {visible_devices}
- Color limits:
  - feature RMS: {figure_info["color_limits"]["feature_rms"]}
  - residual state: {figure_info["color_limits"]["residual_state"]}
  - predicted field: {figure_info["color_limits"]["velocity_field"]}
  - background and final velocity: {figure_info["color_limits"]["background_and_final_velocity"]}
  - final residual: {figure_info["color_limits"]["final_residual"]}

## Exact sampler

The formal model uses identity/pixel-space residual Flow Matching:

    c_R = [C_str; psi_bg(B_hat)]
    R_0 = xi, xi ~ N(0, I)
    t_k = k / 50
    v_k = v_theta(R_k, t_k, c_R)
    R_(k+1) = R_k + (1/50) v_k
    V_hat = B_hat + R_hat

trajectory_arrays.npz stores the real R_k, v_k, Delta R_k, time values,
conditions, background, residual, and final velocity. The stored state arrays
have per-state shape 1 x 70 x 70.

## Validation

- Exact checkpoint load: missing=0, unexpected=0.
- Traced R_50 vs production sample() under the same noise, dtype, and
  condition: torch.testing.assert_close passed.
- Composition V_hat = B_hat + R_hat: torch.testing.assert_close passed.
- All saved arrays are finite.

See trajectory_metadata.json for complete shapes, color limits, hashes, and
validation values. The PPTX keeps labels, equations, and arrows editable; the
spatial tiles are embedded raster data assets generated directly from tensors.
"""
    (output_dir / "README.md").write_text(readme, encoding="utf-8")


def main(argv: list[str] | None = None) -> dict[str, Any]:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--checkpoint", type=Path, default=DEFAULT_CHECKPOINT)
    parser.add_argument("--eval-dir", type=Path, default=DEFAULT_EVAL_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--device", default="cuda:0")
    parser.add_argument("--seed", type=int, default=TRACE_SEED)
    parser.add_argument("--steps", type=int, default=TRACE_STEPS)
    parser.add_argument("--dataset-name", default=TARGET_DATASET)
    parser.add_argument("--source-sample-index", type=int, default=TARGET_SOURCE_SAMPLE_INDEX)
    parser.add_argument("--dataset-id", type=int, default=None)
    parser.add_argument(
        "--official-median",
        action="store_true",
        help="Use the official subset-median held-out record instead of an explicit all-split asset.",
    )
    args = parser.parse_args(argv)
    if args.steps != 50:
        raise ValueError("The formal trajectory protocol requires exactly 50 Euler steps.")
    if not args.config.is_file() or not args.checkpoint.is_file():
        raise FileNotFoundError(f"Missing config or checkpoint: config={args.config}, checkpoint={args.checkpoint}")
    manifest_path = args.eval_dir / "held_out_manifest.csv"
    metrics_path = args.eval_dir / "merged_metrics.csv"
    configure_torch_runtime("medium")
    conf = _prepare_config(args.config)
    explicit_asset = not args.official_median
    dataset_split = "all" if explicit_asset else "test"
    dataset = build_dataset(conf, dataset_split)
    dataset_id = args.dataset_id
    if explicit_asset and dataset_id is None:
        dataset_id = _infer_dataset_id(dataset, args.dataset_name, args.source_sample_index)
    selected = select_requested_record(
        manifest_path,
        metrics_path,
        dataset_name=args.dataset_name,
        source_sample_index=None if args.official_median else args.source_sample_index,
        dataset_id=dataset_id,
    )
    selected_manifest = selected["selected_manifest"]
    selected_metrics = selected["selected_metrics"]
    median_ssim = selected["median_ssim"]
    device = torch.device(args.device)
    if device.type == "cuda" and not torch.cuda.is_available():
        raise RuntimeError(f"Requested {device}, but CUDA is unavailable.")
    trace_result = run_real_trace(
        conf,
        args.checkpoint,
        dataset,
        selected_manifest,
        device=device,
        seed=args.seed,
        steps=args.steps,
    )
    arrays = trace_result["arrays"]
    figure_info = render_figure(arrays, args.output_dir)
    export_editable_pptx(arrays, args.output_dir / "residual_flow_trajectory_editable.pptx")
    write_outputs(
        args.output_dir,
        arrays=arrays,
        trace_result=trace_result,
        selected_manifest=selected_manifest,
        selected_metrics=selected_metrics,
        median_ssim=median_ssim,
        selection_info=selected,
        config_path=args.config,
        checkpoint_path=args.checkpoint,
        eval_dir=args.eval_dir,
        seed=args.seed,
        steps=args.steps,
        figure_info=figure_info,
        visible_devices=os.environ.get("CUDA_VISIBLE_DEVICES"),
    )
    print(
        json.dumps(
            {
                "output_dir": str(args.output_dir),
                "selected_manifest": selected_manifest,
                "selected_metrics": selected_metrics,
                "median_ssim": median_ssim,
                "provenance_scope": selected["provenance_scope"],
                "dataset_split": selected["dataset_split"],
                "validation": trace_result["validation"],
                "files": sorted(path.name for path in args.output_dir.iterdir()),
            },
            indent=2,
        )
    )
    return trace_result


if __name__ == "__main__":
    main()
