from __future__ import annotations

import json
import zipfile

import numpy as np
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from bg_pdr_fm.evaluation.export_condition_contract_pptx import export_condition_contract_pptx
from bg_pdr_fm.evaluation.visualize_condition_contract import render_single_checkpoint_summary_from_csv


def _diagnostic_rows():
    subset_order = ("FlatVelA", "FlatVelB")
    anchor_rows = []
    role_rows = []
    for dataset_id, dataset_name in enumerate(subset_order):
        for sample_index in range(24):
            role_rows.append(
                {
                    "dataset_id": dataset_id,
                    "dataset_name": dataset_name,
                    "source_sample_index": sample_index,
                    "eta_b": 0.72 + 0.01 * dataset_id + 0.03 * np.sin(sample_index),
                    "eta_s": 0.63 + 0.01 * dataset_id + 0.03 * np.cos(sample_index),
                    "chi_bs": 0.16 + 0.02 * dataset_id + 0.04 * abs(np.sin(sample_index)),
                }
            )
            anchor_rows.extend(
                [
                    {
                        "dataset_id": dataset_id,
                        "dataset_name": dataset_name,
                        "source_sample_index": sample_index,
                        "role": "background",
                        "matched": 0.68 + 0.04 * np.sin(sample_index),
                        "shuffled": 0.25 + 0.25 * np.cos(sample_index),
                    },
                    {
                        "dataset_id": dataset_id,
                        "dataset_name": dataset_name,
                        "source_sample_index": sample_index,
                        "role": "structural",
                        "matched": 0.88 + 0.03 * np.cos(sample_index),
                        "shuffled": 0.58 + 0.25 * np.sin(sample_index),
                    },
                ]
            )
    return subset_order, anchor_rows, role_rows


def test_editable_condition_contract_pptx_contains_native_shapes_only(tmp_path):
    subset_order, anchor_rows, role_rows = _diagnostic_rows()
    output = tmp_path / "condition_contract.pptx"
    output_metadata = tmp_path / "condition_contract_pptx_metadata.json"

    export_condition_contract_pptx(
        anchor_rows=anchor_rows,
        role_rows=role_rows,
        output_path=output,
        metadata_output_path=output_metadata,
        subset_order=subset_order,
        expected_count=len(role_rows),
    )

    presentation = Presentation(output)
    assert len(presentation.slides) == 1
    shapes = list(presentation.slides[0].shapes)
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.FREEFORM for shape in shapes) == 4
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes) == 0
    shape_names = {shape.name for shape in shapes}
    assert {
        "Violin_background_matched",
        "Violin_background_shuffled",
        "Violin_structural_matched",
        "Violin_structural_shuffled",
    }.issubset(shape_names)
    assert {f"Role_eta_b_{name}_IQR" for name in subset_order}.issubset(shape_names)
    assert {f"Role_eta_s_{name}_IQR" for name in subset_order}.issubset(shape_names)
    assert {f"Role_chi_bs_{name}_IQR" for name in subset_order}.issubset(shape_names)
    with zipfile.ZipFile(output) as archive:
        assert not any(name.startswith("ppt/media/") for name in archive.namelist())
    report = json.loads(output_metadata.read_text(encoding="utf-8"))
    assert report["all_elements_native_editable"] is True
    assert report["picture_count"] == 0
    assert report["main_figure_panels"] == ["anchor_calibration", "role_specialization"]


def test_two_panel_static_figure_renders_from_existing_csvs(tmp_path):
    subset_order, anchor_rows, role_rows = _diagnostic_rows()
    import csv

    anchor_path = tmp_path / "anchor_calibration.csv"
    with anchor_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(anchor_rows[0]))
        writer.writeheader()
        writer.writerows(anchor_rows)
    role_path = tmp_path / "role_separation.csv"
    with role_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(role_rows[0]))
        writer.writeheader()
        writer.writerows(role_rows)

    outputs = render_single_checkpoint_summary_from_csv(
        output_dir=tmp_path,
        anchor_csv=anchor_path,
        role_csv=role_path,
        dataset_labels=subset_order,
    )

    assert {path.suffix for path in outputs} == {".pdf", ".png", ".svg"}
    from PIL import Image

    with Image.open(tmp_path / "figure3_condition_learning.png") as image:
        assert image.width > image.height
        assert image.info["dpi"][0] >= 300
