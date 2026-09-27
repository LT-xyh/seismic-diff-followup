from __future__ import annotations

import json
import zipfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from bg_pdr_fm.evaluation.export_missing_modality_radar_pptx import (
    export_missing_modality_radar_pptx,
)
from bg_pdr_fm.evaluation.visualize_missing_modality_radar import (
    RADAR_DATASET_ORDER,
    RADAR_METHOD_ORDER,
    RADAR_MODE_ORDER,
)


def _radar_rows() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for method_index, method in enumerate(RADAR_METHOD_ORDER):
        for mode_index, mode in enumerate(RADAR_MODE_ORDER):
            base = 0.58 if mode in {"w/o horizon", "RMS only"} else 0.88
            for dataset_index, dataset_name in enumerate(RADAR_DATASET_ORDER):
                value = base + 0.04 * ((method_index + dataset_index) % 4) / 3.0
                rows.append(
                    {
                        "method": method,
                        "missing_mode": mode,
                        "dataset_name": dataset_name,
                        "num_samples": "1",
                        "ssim": f"{value:.6f}",
                    }
                )
    return rows


def test_editable_radar_pptx_contains_native_shapes_only(tmp_path):
    output = tmp_path / "missing_modality_radar_editable.pptx"
    metadata = tmp_path / "missing_modality_radar_pptx_metadata.json"

    export_missing_modality_radar_pptx(
        rows=_radar_rows(),
        output_path=output,
        metadata_output_path=metadata,
        source_csv=tmp_path / "source.csv",
    )

    presentation = Presentation(output)
    assert len(presentation.slides) == 1
    shapes = list(presentation.slides[0].shapes)
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes) == 0
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.LINE for shape in shapes) > 100
    shape_names = {shape.name for shape in shapes}
    assert "RadarPanel_w/o Well" in shape_names
    assert "RadarCurve_PD-BG-RFM_w/o Well_segment_0" in shape_names
    assert "Legend_PD-BG-RFM" in shape_names

    with zipfile.ZipFile(output) as archive:
        assert not any(name.startswith("ppt/media/") for name in archive.namelist())

    report = json.loads(metadata.read_text(encoding="utf-8"))
    assert report["all_elements_native_editable"] is True
    assert report["picture_count"] == 0
    assert report["num_rows"] == 192
    assert report["panel_radial_ranges"]["w/o well_log"] == [0.85, 1.0]
