from __future__ import annotations

import json
from pathlib import Path

import pytest


ROOT = Path(__file__).resolve().parents[2]
INPUT_DIR = ROOT / "docs" / "paper" / "AAAI2027" / "figures" / "figure1_motivation"


def test_motivation_statistics_pptx_uses_native_editable_shapes(tmp_path: Path) -> None:
    from bg_pdr_fm.evaluation.export_motivation_statistics_pptx import (
        export_motivation_statistics_pptx,
    )

    output = tmp_path / "motivation_statistics_combined_v2_editable.pptx"
    metadata_output = tmp_path / "motivation_statistics_combined_v2_pptx_metadata.json"
    result = export_motivation_statistics_pptx(
        input_dir=INPUT_DIR,
        output_path=output,
        metadata_output_path=metadata_output,
    )

    assert result == output
    assert output.is_file()
    metadata = json.loads(metadata_output.read_text(encoding="utf-8"))
    assert metadata["all_elements_native_editable"] is True
    assert metadata["picture_shape_count"] == 0
    assert metadata["slide_width_in"] / metadata["slide_height_in"] == pytest.approx(504 / 194.4, rel=1e-3)
    assert metadata["native_shape_count"] > 100

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(output)
    assert len(presentation.slides) == 1
    slide = presentation.slides[0]
    assert any(shape.has_text_frame and "Modality" in shape.text for shape in slide.shapes)
    assert any(shape.shape_type == MSO_SHAPE_TYPE.FREEFORM for shape in slide.shapes)
    assert all(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)

