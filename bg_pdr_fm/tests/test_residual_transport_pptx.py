from __future__ import annotations

import json
import zipfile

import numpy as np
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from bg_pdr_fm.evaluation.export_residual_transport_pptx import export_transport_energy_pptx
from bg_pdr_fm.evaluation.visualize_residual_transport import SUBSET_ORDER


def test_editable_transport_pptx_contains_native_shapes_only(tmp_path):
    rows = []
    for dataset_id, name in enumerate(SUBSET_ORDER):
        for source_index, value in enumerate(np.linspace(0.65, 0.92, 20)):
            rows.append(
                {
                    "dataset_id": dataset_id,
                    "dataset_name": name,
                    "source_sample_index": source_index,
                    "q_T": float(value + dataset_id * 0.002),
                }
            )
    values = np.asarray([float(row["q_T"]) for row in rows])
    metadata = {
        "q_t_summary": {
            "n": len(rows),
            "mean_q_t": float(values.mean()),
            "fraction_q_t_lt_1": float(np.mean(values < 1.0)),
            "mean_q_t_ci_low": float(values.mean() - 0.01),
            "mean_q_t_ci_high": float(values.mean() + 0.01),
        }
    }
    output = tmp_path / "transport.pptx"
    output_metadata = tmp_path / "transport_metadata.json"

    export_transport_energy_pptx(
        qt_rows=rows,
        metadata=metadata,
        output_path=output,
        metadata_output_path=output_metadata,
        expected_count=len(rows),
    )

    presentation = Presentation(output)
    assert len(presentation.slides) == 1
    shapes = list(presentation.slides[0].shapes)
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.FREEFORM for shape in shapes) == 8
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes) == 0
    shape_names = {shape.name for shape in shapes}
    assert {f"Violin_{name}" for name in SUBSET_ORDER}.issubset(shape_names)
    assert {f"Box_{name}_IQR" for name in SUBSET_ORDER}.issubset(shape_names)

    with zipfile.ZipFile(output) as archive:
        assert not any(name.startswith("ppt/media/") for name in archive.namelist())
    report = json.loads(output_metadata.read_text(encoding="utf-8"))
    assert report["all_elements_native_editable"] is True
    assert report["picture_count"] == 0
