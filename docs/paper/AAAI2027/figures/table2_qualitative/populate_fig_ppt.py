"""Populate the Table 2 qualitative PowerPoint from the standalone image manifest."""

from __future__ import annotations

import argparse
import csv
import tempfile
from pathlib import Path

import numpy as np
from matplotlib import colormaps
from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


METHODS = (
    "Ground Truth",
    "InversionNet",
    "VelocityGAN",
    "UPFWI",
    "Auto-Linear",
    "Latent U-Net (Large)",
    "PD-BG-RFM",
)
ERROR_METHODS = METHODS[1:]
SLIDE_DATASETS = (
    ("CurveVelA", "CurveVelB", "FlatVelA", "FlatVelB"),
    ("CurveFaultA", "CurveFaultB", "FlatFaultA", "FlatFaultB"),
)


def _clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape.element.getparent().remove(shape.element)


def _add_text(slide, text: str, left: float, top: float, width: float, height: float, size: float, *, vertical: bool = False) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if vertical:
        shape.rotation = 270


def _colorbar(path: Path, colormap: str) -> None:
    values = np.linspace(1.0, 0.0, 512, dtype=np.float32)
    rgba = colormaps[colormap](values, bytes=True)
    pixels = np.repeat(rgba[:, None, :], 28, axis=1)
    Image.fromarray(pixels).save(path)


def _load_images(
    manifest_path: Path,
    image_directory: Path,
    methods: tuple[str, ...],
) -> dict[tuple[str, str], Path]:
    by_key: dict[tuple[str, str], Path] = {}
    with manifest_path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            key = (row["dataset_name"], row["method"])
            if key in by_key:
                raise ValueError(f"Duplicate standalone image manifest entry: {key}.")
            image_path = image_directory / row["filename"]
            if not image_path.is_file():
                raise FileNotFoundError(f"Missing standalone image for {key}: {image_path}.")
            with Image.open(image_path) as image:
                if image.size != (70, 70):
                    raise ValueError(f"Standalone image for {key} must be 70x70, got {image.size}.")
            by_key[key] = image_path
    expected = {(dataset, method) for datasets in SLIDE_DATASETS for dataset in datasets for method in methods}
    if set(by_key) != expected:
        missing = sorted(expected - set(by_key))
        unexpected = sorted(set(by_key) - expected)
        raise ValueError(f"Standalone manifest does not match the PPT grid: missing={missing}, unexpected={unexpected}.")
    return by_key


def _build_grid_slide(
    slide,
    datasets: tuple[str, ...],
    images: dict[tuple[str, str], Path],
    methods: tuple[str, ...],
    colorbar_path: Path,
    colorbar_ticks: tuple[str, str, str, str],
    colorbar_label: str,
    *,
    image_x: float,
) -> int:
    _clear_slide(slide)
    header_y = 0.22
    image_y = 0.72
    row_gap = 0.16
    image_size = 1.28
    column_gap = 0.13
    row_label_x = 0.16
    row_label_width = 1.20
    for column, method in enumerate(methods):
        left = image_x + column * (image_size + column_gap)
        header = "Latent U-Net\n(Large)" if method == "Latent U-Net (Large)" else method
        _add_text(slide, header, left - 0.05, header_y, image_size + 0.10, 0.36, 9.5)

    inserted = 0
    for row, dataset in enumerate(datasets):
        top = image_y + row * (image_size + row_gap)
        _add_text(slide, dataset.replace("A", "\nA").replace("B", "\nB"), row_label_x, top + 0.31, row_label_width, 0.62, 12.0)
        for column, method in enumerate(methods):
            left = image_x + column * (image_size + column_gap)
            slide.shapes.add_picture(str(images[(dataset, method)]), Inches(left), Inches(top), width=Inches(image_size), height=Inches(image_size))
            inserted += 1

    colorbar_x = 11.55
    colorbar_height = len(datasets) * image_size + (len(datasets) - 1) * row_gap
    slide.shapes.add_picture(str(colorbar_path), Inches(colorbar_x), Inches(image_y), width=Inches(0.18), height=Inches(colorbar_height))
    tick_top, tick_upper, tick_lower, tick_bottom = colorbar_ticks
    _add_text(slide, tick_top, colorbar_x + 0.24, image_y - 0.07, 0.55, 0.20, 8.5)
    _add_text(slide, tick_upper, colorbar_x + 0.24, image_y + colorbar_height * 0.22, 0.55, 0.20, 8.5)
    _add_text(slide, tick_lower, colorbar_x + 0.24, image_y + colorbar_height * 0.70, 0.55, 0.20, 8.5)
    _add_text(slide, tick_bottom, colorbar_x + 0.24, image_y + colorbar_height - 0.13, 0.55, 0.20, 8.5)
    _add_text(slide, colorbar_label, colorbar_x + 0.71, image_y + colorbar_height / 2.0 - 0.55, 0.24, 1.10, 8.5, vertical=True)
    return inserted


def populate(
    template: Path,
    manifest: Path,
    image_directory: Path,
    output: Path,
    *,
    error_manifest: Path,
    error_image_directory: Path,
) -> int:
    velocity_images = _load_images(manifest, image_directory, METHODS)
    error_images = _load_images(error_manifest, error_image_directory, ERROR_METHODS)
    presentation = Presentation(template)
    if not presentation.slides:
        raise ValueError(f"Template has no slides: {template}.")
    first_slide = presentation.slides[0]
    for extra_slide in list(presentation.slides)[1:]:
        slide_id = next(
            item
            for item in list(presentation.slides._sldIdLst)
            if presentation.part.related_part(item.rId) is extra_slide.part
        )
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
    second_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    third_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    fourth_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    with tempfile.TemporaryDirectory() as directory:
        colorbar_path = Path(directory) / "jet_velocity_colorbar.png"
        error_colorbar_path = Path(directory) / "inferno_error_colorbar.png"
        _colorbar(colorbar_path, "jet")
        _colorbar(error_colorbar_path, "inferno")
        inserted = _build_grid_slide(
            first_slide,
            SLIDE_DATASETS[0],
            velocity_images,
            METHODS,
            colorbar_path,
            ("4500", "3750", "2250", "1500"),
            "Velocity (m/s)",
            image_x=1.55,
        )
        inserted += _build_grid_slide(
            second_slide,
            SLIDE_DATASETS[1],
            velocity_images,
            METHODS,
            colorbar_path,
            ("4500", "3750", "2250", "1500"),
            "Velocity (m/s)",
            image_x=1.55,
        )
        inserted += _build_grid_slide(
            third_slide,
            SLIDE_DATASETS[0],
            error_images,
            ERROR_METHODS,
            error_colorbar_path,
            ("750", "562.5", "187.5", "0"),
            "Absolute error (m/s)",
            image_x=2.25,
        )
        inserted += _build_grid_slide(
            fourth_slide,
            SLIDE_DATASETS[1],
            error_images,
            ERROR_METHODS,
            error_colorbar_path,
            ("750", "562.5", "187.5", "0"),
            "Absolute error (m/s)",
            image_x=2.25,
        )
        output.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(output)
    return inserted


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--images", type=Path, required=True)
    parser.add_argument("--error-manifest", type=Path, required=True)
    parser.add_argument("--errors", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    inserted = populate(
        args.template,
        args.manifest,
        args.images,
        args.output,
        error_manifest=args.error_manifest,
        error_image_directory=args.errors,
    )
    if inserted != 104:
        raise RuntimeError(f"Expected 104 inserted images, got {inserted}.")
    print(f"Inserted {inserted} images into {args.output}.")


if __name__ == "__main__":
    main()
