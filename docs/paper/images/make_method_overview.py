"""Create the editable PD-BG-RFM method overview figure.

The PowerPoint contains native text, boxes, lines, and arrows. Scientific
image tiles remain pictures because they are data assets. An SVG is generated
from the same layout and converted to PDF/PNG for manuscript export.
"""

from __future__ import annotations

import argparse
import base64
import html
import math
import json
import subprocess
import zipfile
from pathlib import Path
from typing import Iterable

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


REPO_ROOT = Path(__file__).resolve().parents[3]
ASSET_ROOT = REPO_ROOT / "docs" / "paper" / "images" / "figure2_assets"
SOURCE_PPTX = REPO_ROOT / "docs" / "paper" / "AAAI2027" / "figures" / "method.pptx"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "docs" / "paper" / "AAAI2027" / "figures"

SLIDE_W = 13.333
SLIDE_H = 5.0
SVG_SCALE = 100.0

GREEN = "2A9D6F"
GREEN_LIGHT = "EAF6F0"
BLUE = "246BCE"
BLUE_LIGHT = "EAF2FC"
ORANGE = "D9822B"
ORANGE_LIGHT = "FFF3E7"
PURPLE = "7257C7"
PURPLE_LIGHT = "F1EDFC"
GRAY = "667085"
GRAY_LIGHT = "F2F4F7"
PANEL_LINE = "C8D1DC"
TEXT = "24303D"
MUTED = "5E6B7A"
ARROW = "344054"
WHITE = "FFFFFF"


TILES = {
    "pstm": ASSET_ROOT / "data_tiles" / "pstm_1to1.png",
    "horizon": ASSET_ROOT / "data_tiles" / "horizon_1to1.png",
    "rms": ASSET_ROOT / "data_tiles" / "rms_1to1.png",
    "well": ASSET_ROOT / "data_tiles" / "well_1to1.png",
    "velocity": ASSET_ROOT / "data_tiles" / "velocity_1to1.png",
    "background": ASSET_ROOT / "data_tiles" / "background_1to1.png",
    "structure": ASSET_ROOT / "data_tiles" / "structure_1to1.png",
    "noise": ASSET_ROOT / "data_tiles" / "noise_field.png",
    "rt": ASSET_ROOT / "data_tiles" / "interpolated_residual_field.png",
    "rhat": ASSET_ROOT / "transport_burden" / "subfigures" / "stage3_target_residual_Rhat.png",
    "cb": ASSET_ROOT / "feature_maps" / "cb_aggregate.png",
    "cs": ASSET_ROOT / "feature_maps" / "cs_aggregate.png",
}


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _pt(value: float) -> float:
    return value * 1.3333333333


def _svg_color(value: str) -> str:
    return f"#{value}"


def _asset_data_uri(path: Path) -> str:
    mime = "image/png" if path.suffix.lower() == ".png" else "image/jpeg"
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime};base64,{encoded}"


def _check_assets() -> None:
    if not SOURCE_PPTX.is_file():
        raise FileNotFoundError(f"Source asset deck is missing: {SOURCE_PPTX}")
    missing = [str(path) for path in TILES.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing method overview assets:\n" + "\n".join(missing))


def _add_text(
    slide,
    text: str,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    size: float = 12.0,
    color: str = TEXT,
    bold: bool = False,
    italic: bool = False,
    align: str = "left",
    font: str = "Arial",
    name: str | None = None,
):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)
    return shape


def _add_box(
    slide,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    fill: str = WHITE,
    line: str = PANEL_LINE,
    radius: bool = True,
    dashed: bool = False,
    name: str,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.color.rgb = _rgb(line)
    shape.line.width = Pt(0.8)
    if dashed:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shape


def _add_picture(slide, path: Path, *, left: float, top: float, width: float, height: float, name: str):
    shape = slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))
    shape.name = name
    return shape


def _add_line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str = ARROW,
    width: float = 1.1,
    dashed: bool = False,
    name: str,
    arrowhead: bool = True,
):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.name = name
    line.line.color.rgb = _rgb(color)
    line.line.width = Pt(width)
    if dashed:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrowhead:
        # python-pptx 1.0 does not expose endArrow; use a native triangle.
        head = slide.shapes.add_shape(
            MSO_SHAPE.ISOSCELES_TRIANGLE,
            Inches(x2 - 0.07),
            Inches(y2 - 0.055),
            Inches(0.14),
            Inches(0.11),
        )
        head.rotation = 90 + math.degrees(math.atan2(y2 - y1, x2 - x1))
        head.name = f"{name}_ArrowHead"
        head.fill.solid()
        head.fill.fore_color.rgb = _rgb(color)
        head.line.color.rgb = _rgb(color)
    return line


def _add_svg_text(parts: list[str], text: str, *, x: float, y: float, w: float, h: float, size: float,
                  color: str = TEXT, bold: bool = False, italic: bool = False, align: str = "left",
                  font: str = "Arial", name: str | None = None) -> None:
    anchor = {"left": "start", "center": "middle", "right": "end"}[align]
    tx = x if align == "left" else x + (w / 2 if align == "center" else w)
    ty = y + h * 0.72
    attrs = [
        f'font-family="{html.escape(font)}"',
        f'font-size="{_pt(size):.2f}px"',
        f'fill="{_svg_color(color)}"',
        f'text-anchor="{anchor}"',
    ]
    if bold:
        attrs.append('font-weight="700"')
    if italic:
        attrs.append('font-style="italic"')
    if name:
        attrs.append(f'id="{html.escape(name)}"')
    lines = text.splitlines() or [""]
    if len(lines) == 1:
        attrs.extend([f'x="{tx * SVG_SCALE:.2f}"', f'y="{ty * SVG_SCALE:.2f}"'])
        parts.append(f"<text {' '.join(attrs)}>{html.escape(lines[0])}</text>")
        return
    line_height = _pt(size) * 1.12
    first_y = ty * SVG_SCALE - ((len(lines) - 1) * line_height / 2.0)
    attrs.extend([f'x="{tx * SVG_SCALE:.2f}"', f'y="{first_y:.2f}"'])
    tspans = []
    for index, line in enumerate(lines):
        dy = "0" if index == 0 else f"{line_height:.2f}px"
        tspans.append(f'<tspan x="{tx * SVG_SCALE:.2f}" dy="{dy}">{html.escape(line)}</tspan>')
    parts.append(f"<text {' '.join(attrs)}>{''.join(tspans)}</text>")


def _add_svg_box(parts: list[str], *, left: float, top: float, width: float, height: float,
                 fill: str = WHITE, line: str = PANEL_LINE, radius: float = 0.08,
                 dashed: bool = False, name: str) -> None:
    dash = ' stroke-dasharray="7,5"' if dashed else ""
    parts.append(
        f'<rect id="{html.escape(name)}" x="{left * SVG_SCALE:.2f}" y="{top * SVG_SCALE:.2f}" '
        f'width="{width * SVG_SCALE:.2f}" height="{height * SVG_SCALE:.2f}" '
        f'rx="{radius * SVG_SCALE:.2f}" fill="{_svg_color(fill)}" '
        f'stroke="{_svg_color(line)}" stroke-width="1.2"{dash}/>'
    )


def _add_svg_image(parts: list[str], path: Path, *, left: float, top: float, width: float, height: float,
                   name: str) -> None:
    parts.append(
        f'<image id="{html.escape(name)}" x="{left * SVG_SCALE:.2f}" y="{top * SVG_SCALE:.2f}" '
        f'width="{width * SVG_SCALE:.2f}" height="{height * SVG_SCALE:.2f}" '
        f'preserveAspectRatio="none" href="{_asset_data_uri(path)}" '
        f'xlink:href="{_asset_data_uri(path)}"/>'
    )


def _add_svg_line(parts: list[str], x1: float, y1: float, x2: float, y2: float, *, color: str = ARROW,
                  width: float = 1.3, dashed: bool = False, arrow: bool = True, name: str) -> None:
    dash = ' stroke-dasharray="7,5"' if dashed else ""
    marker = ' marker-end="url(#arrow)"' if arrow and not dashed else (' marker-end="url(#orange-arrow)"' if arrow else "")
    parts.append(
        f'<line id="{html.escape(name)}" x1="{x1 * SVG_SCALE:.2f}" y1="{y1 * SVG_SCALE:.2f}" '
        f'x2="{x2 * SVG_SCALE:.2f}" y2="{y2 * SVG_SCALE:.2f}" stroke="{_svg_color(color)}" '
        f'stroke-width="{width * 1.333:.2f}" stroke-linecap="round"{dash}{marker}/>'
    )


def _draw_ppt(slide) -> None:
    # Canvas legend.
    _add_text(slide, "Solid arrows: inference computation", left=0.22, top=0.10, width=2.35, height=0.20,
              size=8.5, color=MUTED, name="Legend_Inference")
    _add_line(slide, 2.25, 0.20, 2.72, 0.20, color=ARROW, width=0.9, name="Legend_Inference_Line")
    _add_text(slide, "Dashed orange: training-only supervision", left=2.82, top=0.10, width=2.75, height=0.20,
              size=8.5, color=ORANGE, name="Legend_Training")
    _add_line(slide, 5.47, 0.20, 5.95, 0.20, color=ORANGE, width=0.9, dashed=True,
              name="Legend_Training_Line", arrowhead=False)

    # Main panels.
    _add_box(slide, left=0.18, top=0.42, width=5.25, height=4.35, fill="FBFCFE", line=PANEL_LINE,
             radius=True, name="Panel_A_Condition_Contract")
    _add_box(slide, left=5.58, top=0.42, width=7.56, height=4.35, fill="FBFCFE", line=PANEL_LINE,
             radius=True, name="Panel_B_Residual_Transport")
    _add_text(slide, "(a) Physics-Decoupled Condition Contract", left=0.38, top=0.52, width=4.85, height=0.27,
              size=14.5, color=TEXT, bold=True, align="center", name="Panel_A_Title")
    _add_text(slide, "(b) Prediction-Consistent Residual Transport", left=5.82, top=0.52, width=7.08, height=0.27,
              size=14.5, color=TEXT, bold=True, align="center", name="Panel_B_Title")

    # Compact factorization formula.
    _add_box(slide, left=2.00, top=0.80, width=3.10, height=0.42, fill=WHITE, line=PANEL_LINE,
             radius=False, name="Factorization_Formula_Box")
    _add_text(slide, "p_PD(b,r|XΩ) = δ[b=f_B(C_B)] pθ(r|C_S,b)", left=2.08, top=0.83, width=2.94, height=0.17,
              size=9.1, color=PURPLE, bold=True, align="center", font="Cambria Math", name="Factorization_Formula")
    _add_text(slide, "deterministic background + distributional residual", left=2.08, top=1.04, width=2.94, height=0.12,
              size=7.2, color=MUTED, italic=True, align="center", name="Factorization_Label")

    # Four observed inputs.
    input_specs = [
        ("PoSTM", TILES["pstm"], 0.40, 1.35, "Input_PoSTM"),
        ("Horizon", TILES["horizon"], 1.18, 1.35, "Input_Horizon"),
        ("RMS", TILES["rms"], 0.40, 2.48, "Input_RMS"),
        ("Well", TILES["well"], 1.18, 2.48, "Input_Well"),
    ]
    for label, path, left, top, name in input_specs:
        _add_picture(slide, path, left=left, top=top, width=0.63, height=0.63, name=f"{name}_Tile")
        _add_text(slide, label, left=left - 0.06, top=top + 0.67, width=0.75, height=0.15, size=8.3,
                  color=TEXT, bold=True, align="center", name=f"{name}_Label")
    _add_text(slide, "asymmetric evidence", left=0.38, top=3.44, width=1.92, height=0.20, size=8.8,
              color=MUTED, italic=True, align="center", name="Input_Asymmetry_Label")

    # Encoder and role heads.
    _add_box(slide, left=2.05, top=1.45, width=0.88, height=1.62, fill=WHITE, line=PANEL_LINE,
             radius=True, name="Modality_Evidence_Encoder")
    _add_text(slide, "Modality\nEvidence\nEncoder", left=2.13, top=1.73, width=0.72, height=0.62, size=9.0,
              color=TEXT, bold=True, align="center", name="Modality_Evidence_Encoder_Label")
    _add_text(slide, "E_m", left=2.27, top=2.55, width=0.45, height=0.18, size=10.5, color=BLUE,
              bold=True, align="center", font="Cambria Math", name="Encoder_Em")
    for idx, (_, _, left, top, name) in enumerate(input_specs):
        _add_line(slide, left + 0.63, top + 0.32, 2.05, 1.80 + idx * 0.17, color=ARROW, width=0.65,
                  name=f"{name}_to_Encoder", arrowhead=False)

    _add_box(slide, left=3.15, top=1.42, width=0.90, height=1.70, fill=GRAY_LIGHT, line=PANEL_LINE,
             radius=True, name="Role_Heads")
    _add_text(slide, "role heads", left=3.24, top=1.50, width=0.72, height=0.16, size=8.4,
              color=TEXT, bold=True, align="center", name="Role_Heads_Title")
    _add_text(slide, "B_m", left=3.31, top=1.88, width=0.58, height=0.20, size=10.7, color=GREEN,
              bold=True, align="center", font="Cambria Math", name="Role_Head_B")
    _add_text(slide, "S_m", left=3.31, top=2.23, width=0.58, height=0.20, size=10.7, color=BLUE,
              bold=True, align="center", font="Cambria Math", name="Role_Head_S")
    _add_text(slide, "U_m", left=3.31, top=2.58, width=0.58, height=0.20, size=10.7, color=GRAY,
              bold=True, align="center", font="Cambria Math", name="Role_Head_U")
    _add_line(slide, 2.93, 2.25, 3.15, 2.25, color=ARROW, width=1.0, name="Encoder_to_RoleHeads")

    _add_box(slide, left=4.24, top=1.60, width=0.82, height=0.54, fill=GREEN_LIGHT, line=GREEN,
             radius=True, name="CB_Fusion")
    _add_text(slide, "C_B", left=4.31, top=1.73, width=0.68, height=0.20, size=13.0, color=GREEN,
              bold=True, align="center", font="Cambria Math", name="CB_Label")
    _add_box(slide, left=4.24, top=2.33, width=0.82, height=0.54, fill=BLUE_LIGHT, line=BLUE,
             radius=True, name="CS_Fusion")
    _add_text(slide, "C_S", left=4.31, top=2.46, width=0.68, height=0.20, size=13.0, color=BLUE,
              bold=True, align="center", font="Cambria Math", name="CS_Label")
    _add_line(slide, 4.05, 1.98, 4.24, 1.87, color=GREEN, width=1.0, name="Bm_to_CB")
    _add_line(slide, 4.05, 2.34, 4.24, 2.60, color=BLUE, width=1.0, name="Sm_to_CS")
    _add_line(slide, 4.05, 2.75, 4.24, 2.75, color=GRAY, width=0.8, name="Um_Separation_End", arrowhead=False)
    _add_text(slide, "reliability\nfusion", left=4.03, top=3.02, width=1.02, height=0.28, size=7.3,
              color=MUTED, italic=True, align="center", name="Reliability_Fusion_Label")
    _add_text(slide, "U_m: separation only", left=3.10, top=3.22, width=1.75, height=0.16, size=7.3,
              color=GRAY, italic=True, align="center", name="Unique_Head_Label")

    # Actual feature maps make C_B/C_S semantic without turning the figure into a diagnostic panel.
    _add_picture(slide, TILES["cb"], left=4.16, top=3.56, width=0.47, height=0.47, name="CB_Feature_Map")
    _add_picture(slide, TILES["cs"], left=4.75, top=3.56, width=0.47, height=0.47, name="CS_Feature_Map")
    _add_text(slide, "learned role maps", left=4.13, top=4.08, width=1.13, height=0.14, size=7.1,
              color=MUTED, italic=True, align="center", name="Role_Maps_Label")

    # Training-only anchors.
    _add_box(slide, left=2.02, top=3.62, width=1.86, height=0.64, fill=ORANGE_LIGHT, line=ORANGE,
             radius=True, dashed=True, name="Training_Only_Anchors")
    _add_text(slide, "training-only anchors", left=2.13, top=3.73, width=1.64, height=0.15, size=8.2,
              color=ORANGE, bold=True, align="center", name="Training_Only_Anchors_Title")
    _add_text(slide, "A_B=W_L(V),  A_S=W_H(V)", left=2.10, top=3.96, width=1.70, height=0.17, size=8.4,
              color=ORANGE, align="center", font="Cambria Math", name="Training_Only_Anchors_Formula")
    _add_text(slide, "not available at inference", left=2.13, top=4.18, width=1.64, height=0.13, size=7.1,
              color=MUTED, italic=True, align="center", name="Training_Only_Anchors_Note")
    _add_line(slide, 2.95, 3.62, 3.58, 3.12, color=ORANGE, width=0.85, dashed=True,
              name="Anchor_to_RoleHeads", arrowhead=False)
    _add_line(slide, 3.58, 3.12, 4.24, 2.60, color=ORANGE, width=0.85, dashed=True,
              name="Anchor_to_CS", arrowhead=False)

    # Panel B: directed responsibility assignment.
    _add_text(slide, "restricted routing", left=5.87, top=0.86, width=1.25, height=0.16, size=8.1,
              color=MUTED, italic=True, name="Restricted_Routing_Label")
    _add_box(slide, left=5.92, top=1.25, width=0.76, height=0.48, fill=GREEN_LIGHT, line=GREEN,
             radius=True, name="B_Condition_Node")
    _add_text(slide, "C_B", left=5.99, top=1.38, width=0.62, height=0.18, size=12.0, color=GREEN,
              bold=True, align="center", font="Cambria Math", name="B_Condition_Label")
    _add_box(slide, left=6.88, top=1.12, width=1.42, height=0.74, fill=GREEN_LIGHT, line=GREEN,
             radius=True, name="Background_UNet")
    _add_text(slide, "Background U-Net", left=7.00, top=1.26, width=1.18, height=0.17, size=9.0,
              color=GREEN, bold=True, align="center", name="Background_UNet_Title")
    _add_text(slide, "f_B(C_B)", left=7.13, top=1.53, width=0.92, height=0.17, size=9.2,
              color=GREEN, align="center", font="Cambria Math", name="Background_UNet_Formula")
    _add_box(slide, left=8.52, top=1.25, width=0.85, height=0.48, fill=GREEN_LIGHT, line=GREEN,
             radius=True, name="Predicted_Background")
    _add_text(slide, "B̂", left=8.58, top=1.38, width=0.73, height=0.18, size=11.5, color=GREEN,
              bold=True, align="center", font="Cambria Math", name="Predicted_Background_Label")
    _add_line(slide, 6.68, 1.49, 6.88, 1.49, color=GREEN, width=1.2, name="CB_to_Background")
    _add_line(slide, 8.30, 1.49, 8.52, 1.49, color=GREEN, width=1.2, name="Background_to_Bhat")

    _add_box(slide, left=8.45, top=1.98, width=1.02, height=0.48, fill=WHITE, line=GREEN,
             radius=True, name="Stop_Gradient_Node")
    _add_text(slide, "B̄=sg[B̂]", left=8.50, top=2.11, width=0.92, height=0.18, size=8.0,
              color=GREEN, bold=True, align="center", font="Cambria Math", name="Stop_Gradient_Label")
    _add_line(slide, 8.95, 1.73, 8.95, 1.98, color=GREEN, width=1.0, name="Bhat_to_StopGradient")
    _add_text(slide, "∇ωB L_R = 0", left=9.55, top=1.99, width=1.08, height=0.18, size=7.7,
              color=PURPLE, italic=True, font="Cambria Math", name="Stop_Gradient_Constraint")

    _add_box(slide, left=5.92, top=2.52, width=0.76, height=0.48, fill=BLUE_LIGHT, line=BLUE,
             radius=True, name="S_Condition_Node")
    _add_text(slide, "C_S", left=5.99, top=2.65, width=0.62, height=0.18, size=12.0, color=BLUE,
              bold=True, align="center", font="Cambria Math", name="S_Condition_Label")
    _add_line(slide, 5.06, 1.87, 5.92, 1.49, color=GREEN, width=1.0, name="Contract_CB_to_Responsibility")
    _add_line(slide, 5.06, 2.60, 5.92, 2.76, color=BLUE, width=1.0, name="Contract_CS_to_Responsibility")
    _add_box(slide, left=9.67, top=2.46, width=1.35, height=0.62, fill=BLUE_LIGHT, line=BLUE,
             radius=True, name="Residual_Context")
    _add_text(slide, "c=[C_S;ψ_B(B̄)]", left=9.76, top=2.61, width=1.17, height=0.18, size=8.1,
              color=BLUE, bold=True, align="center", font="Cambria Math", name="Residual_Context_Formula")
    _add_line(slide, 8.95, 2.46, 9.67, 2.77, color=GREEN, width=1.0, name="Bbar_to_Context")
    _add_line(slide, 6.68, 2.76, 9.67, 2.76, color=BLUE, width=1.1, name="CS_to_Context")

    _add_box(slide, left=10.13, top=3.24, width=2.05, height=0.75, fill=BLUE_LIGHT, line=BLUE,
             radius=True, name="Residual_FM_Block")
    _add_text(slide, "velocity-space FiLM", left=10.26, top=3.36, width=1.80, height=0.16, size=9.1,
              color=BLUE, bold=True, align="center", name="Residual_FM_Title")
    _add_text(slide, "Residual Flow Matching", left=10.28, top=3.60, width=1.76, height=0.16, size=8.7,
              color=BLUE, align="center", name="Residual_FM_Subtitle")
    _add_line(slide, 10.35, 3.08, 10.35, 3.24, color=BLUE, width=1.1, name="Context_to_FM")

    # Three-stage transport tiles.
    transport_specs = [
        ("noise", "ξ", 8.00, 3.39, "Transport_Noise"),
        ("rt", "R_t", 8.65, 3.39, "Transport_Rt"),
        ("rhat", "R_{B̂}", 9.30, 3.39, "Transport_Rhat"),
    ]
    for key, label, left, top, name in transport_specs:
        _add_picture(slide, TILES[key], left=left, top=top, width=0.48, height=0.48, name=f"{name}_Tile")
        _add_text(slide, label, left=left - 0.02, top=top + 0.50, width=0.52, height=0.14, size=7.4,
                  color=BLUE, align="center", font="Cambria Math", name=f"{name}_Label")
    _add_line(slide, 8.49, 3.63, 8.65, 3.63, color=BLUE, width=0.8, name="Noise_to_Rt")
    _add_line(slide, 9.14, 3.63, 9.30, 3.63, color=BLUE, width=0.8, name="Rt_to_Rhat")
    _add_text(slide, "R_t=(1-t)ξ+tR_{B̂}", left=7.94, top=4.09, width=1.90, height=0.17, size=8.0,
              color=BLUE, align="center", font="Cambria Math", name="Transport_Interpolation_Formula")
    _add_text(slide, "u_t=R_{B̂}-ξ", left=9.72, top=4.09, width=1.10, height=0.17, size=8.0,
              color=BLUE, align="center", font="Cambria Math", name="Transport_Target_Formula")

    # Training-only residual target enters the residual transport.
    _add_box(slide, left=6.00, top=3.76, width=1.68, height=0.55, fill=ORANGE_LIGHT, line=ORANGE,
             radius=True, dashed=True, name="Residual_Target_TrainingOnly")
    _add_text(slide, "V (training only)", left=6.10, top=3.86, width=1.48, height=0.15, size=8.0,
              color=ORANGE, bold=True, align="center", name="Residual_Target_V_Label")
    _add_text(slide, "R_{B̂}=V-B̄", left=6.12, top=4.08, width=1.44, height=0.16, size=8.3,
              color=ORANGE, align="center", font="Cambria Math", name="Residual_Target_Formula")
    _add_line(slide, 7.68, 4.03, 9.30, 4.03, color=ORANGE, width=0.85, dashed=True,
              name="Target_to_Residual_Transport", arrowhead=False)

    # Composition endpoint.
    _add_box(slide, left=11.35, top=1.36, width=0.43, height=0.43, fill=WHITE, line=TEXT,
             radius=False, name="Composition_Plus")
    _add_text(slide, "+", left=11.35, top=1.42, width=0.43, height=0.20, size=15.0, color=TEXT,
              bold=True, align="center", name="Composition_Plus_Label")
    _add_line(slide, 9.37, 1.49, 11.35, 1.58, color=GREEN, width=1.1, name="Bhat_to_Composition")
    _add_line(slide, 11.13, 3.60, 11.55, 1.79, color=BLUE, width=1.1, name="Rhat_to_Composition")
    _add_box(slide, left=11.82, top=1.24, width=1.18, height=0.68, fill=WHITE, line=TEXT,
             radius=True, name="Final_Output")
    _add_text(slide, "V̂ = B̂ + R̂", left=11.90, top=1.46, width=1.02, height=0.22, size=8.2,
              color=TEXT, bold=True, align="center", font="Cambria Math", name="Final_Output_Formula")
    _add_line(slide, 11.78, 1.58, 11.82, 1.58, color=TEXT, width=1.0, name="Composition_to_Output")
    _add_text(slide, "V̂-V=R̂-R_{B̂}", left=11.16, top=2.04, width=1.75, height=0.18, size=8.0,
              color=PURPLE, align="center", font="Cambria Math", name="Composition_Consistency")

    _add_box(slide, left=10.72, top=4.14, width=2.35, height=0.56, fill=PURPLE_LIGHT, line=PURPLE,
             radius=True, name="Transport_Energy_Condition")
    _add_text(slide, "E||u_R||² < E||u_V||²", left=10.84, top=4.20, width=2.10, height=0.16,
              size=10.4, color=PURPLE, align="center", font="Cambria Math", name="Transport_Energy_Formula_1")
    _add_text(slide, "iff E||V-B̂||² < E||V||²", left=10.84, top=4.39, width=2.10, height=0.16,
              size=10.4, color=PURPLE, align="center", font="Cambria Math", name="Transport_Energy_Formula_2")


def _svg_header() -> list[str]:
    return [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
        f'width="{SLIDE_W}in" height="{SLIDE_H}in" viewBox="0 0 {SLIDE_W * SVG_SCALE:.2f} {SLIDE_H * SVG_SCALE:.2f}">',
        '<defs>',
        f'<marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8 z" fill="{_svg_color(ARROW)}"/></marker>',
        f'<marker id="orange-arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8 z" fill="{_svg_color(ORANGE)}"/></marker>',
        '</defs>',
        f'<rect x="0" y="0" width="{SLIDE_W * SVG_SCALE:.2f}" height="{SLIDE_H * SVG_SCALE:.2f}" fill="#FFFFFF"/>',
    ]


def _draw_svg() -> str:
    parts = _svg_header()
    _add_svg_text(parts, "Solid arrows: inference computation", x=0.22, y=0.10, w=2.35, h=0.20, size=8.5, color=MUTED, name="Legend_Inference")
    _add_svg_line(parts, 2.25, 0.20, 2.72, 0.20, color=ARROW, width=0.9, name="Legend_Inference_Line")
    _add_svg_text(parts, "Dashed orange: training-only supervision", x=2.82, y=0.10, w=2.75, h=0.20, size=8.5, color=ORANGE, name="Legend_Training")
    _add_svg_line(parts, 5.47, 0.20, 5.95, 0.20, color=ORANGE, width=0.9, dashed=True, arrow=False, name="Legend_Training_Line")

    _add_svg_box(parts, left=0.18, top=0.42, width=5.25, height=4.35, fill="FBFCFE", line=PANEL_LINE, name="Panel_A_Condition_Contract")
    _add_svg_box(parts, left=5.58, top=0.42, width=7.56, height=4.35, fill="FBFCFE", line=PANEL_LINE, name="Panel_B_Residual_Transport")
    _add_svg_text(parts, "(a) Physics-Decoupled Condition Contract", x=0.38, y=0.52, w=4.85, h=0.27, size=14.5, bold=True, align="center", name="Panel_A_Title")
    _add_svg_text(parts, "(b) Prediction-Consistent Residual Transport", x=5.82, y=0.52, w=7.08, h=0.27, size=14.5, bold=True, align="center", name="Panel_B_Title")

    _add_svg_box(parts, left=2.00, top=0.80, width=3.10, height=0.42, fill=WHITE, line=PANEL_LINE, radius=0, name="Factorization_Formula_Box")
    _add_svg_text(parts, "p_PD(b,r|XΩ) = δ[b=f_B(C_B)] pθ(r|C_S,b)", x=2.08, y=0.83, w=2.94, h=0.17, size=9.1, color=PURPLE, bold=True, align="center", font="Cambria Math", name="Factorization_Formula")
    _add_svg_text(parts, "deterministic background + distributional residual", x=2.08, y=1.04, w=2.94, h=0.12, size=7.2, color=MUTED, italic=True, align="center", name="Factorization_Label")

    input_specs = [
        ("PoSTM", TILES["pstm"], 0.40, 1.35, "Input_PoSTM"),
        ("Horizon", TILES["horizon"], 1.18, 1.35, "Input_Horizon"),
        ("RMS", TILES["rms"], 0.40, 2.48, "Input_RMS"),
        ("Well", TILES["well"], 1.18, 2.48, "Input_Well"),
    ]
    for label, path, left, top, name in input_specs:
        _add_svg_image(parts, path, left=left, top=top, width=0.63, height=0.63, name=f"{name}_Tile")
        _add_svg_text(parts, label, x=left - 0.06, y=top + 0.67, w=0.75, h=0.15, size=8.3, bold=True, align="center", name=f"{name}_Label")
    _add_svg_text(parts, "asymmetric evidence", x=0.38, y=3.44, w=1.92, h=0.20, size=8.8, color=MUTED, italic=True, align="center", name="Input_Asymmetry_Label")

    _add_svg_box(parts, left=2.05, top=1.45, width=0.88, height=1.62, fill=WHITE, line=PANEL_LINE, name="Modality_Evidence_Encoder")
    _add_svg_text(parts, "Modality\nEvidence\nEncoder", x=2.13, y=1.73, w=0.72, h=0.62, size=9.0, bold=True, align="center", name="Modality_Evidence_Encoder_Label")
    _add_svg_text(parts, "E_m", x=2.27, y=2.55, w=0.45, h=0.18, size=10.5, color=BLUE, bold=True, align="center", font="Cambria Math", name="Encoder_Em")
    for idx, (_, _, left, top, name) in enumerate(input_specs):
        _add_svg_line(parts, left + 0.63, top + 0.32, 2.05, 1.80 + idx * 0.17, color=ARROW, width=0.65, arrow=False, name=f"{name}_to_Encoder")
    _add_svg_box(parts, left=3.15, top=1.42, width=0.90, height=1.70, fill=GRAY_LIGHT, line=PANEL_LINE, name="Role_Heads")
    _add_svg_text(parts, "role heads", x=3.24, y=1.50, w=0.72, h=0.16, size=8.4, bold=True, align="center", name="Role_Heads_Title")
    _add_svg_text(parts, "B_m", x=3.31, y=1.88, w=0.58, h=0.20, size=10.7, color=GREEN, bold=True, align="center", font="Cambria Math", name="Role_Head_B")
    _add_svg_text(parts, "S_m", x=3.31, y=2.23, w=0.58, h=0.20, size=10.7, color=BLUE, bold=True, align="center", font="Cambria Math", name="Role_Head_S")
    _add_svg_text(parts, "U_m", x=3.31, y=2.58, w=0.58, h=0.20, size=10.7, color=GRAY, bold=True, align="center", font="Cambria Math", name="Role_Head_U")
    _add_svg_line(parts, 2.93, 2.25, 3.15, 2.25, color=ARROW, width=1.0, name="Encoder_to_RoleHeads")
    _add_svg_box(parts, left=4.24, top=1.60, width=0.82, height=0.54, fill=GREEN_LIGHT, line=GREEN, name="CB_Fusion")
    _add_svg_text(parts, "C_B", x=4.31, y=1.73, w=0.68, h=0.20, size=13.0, color=GREEN, bold=True, align="center", font="Cambria Math", name="CB_Label")
    _add_svg_box(parts, left=4.24, top=2.33, width=0.82, height=0.54, fill=BLUE_LIGHT, line=BLUE, name="CS_Fusion")
    _add_svg_text(parts, "C_S", x=4.31, y=2.46, w=0.68, h=0.20, size=13.0, color=BLUE, bold=True, align="center", font="Cambria Math", name="CS_Label")
    _add_svg_line(parts, 4.05, 1.98, 4.24, 1.87, color=GREEN, width=1.0, name="Bm_to_CB")
    _add_svg_line(parts, 4.05, 2.34, 4.24, 2.60, color=BLUE, width=1.0, name="Sm_to_CS")
    _add_svg_line(parts, 4.05, 2.75, 4.24, 2.75, color=GRAY, width=0.8, arrow=False, name="Um_Separation_End")
    _add_svg_text(parts, "reliability\nfusion", x=4.03, y=3.02, w=1.02, h=0.28, size=7.3, color=MUTED, italic=True, align="center", name="Reliability_Fusion_Label")
    _add_svg_text(parts, "U_m: separation only", x=3.10, y=3.22, w=1.75, h=0.16, size=7.3, color=GRAY, italic=True, align="center", name="Unique_Head_Label")
    _add_svg_image(parts, TILES["cb"], left=4.16, top=3.56, width=0.47, height=0.47, name="CB_Feature_Map")
    _add_svg_image(parts, TILES["cs"], left=4.75, top=3.56, width=0.47, height=0.47, name="CS_Feature_Map")
    _add_svg_text(parts, "learned role maps", x=4.13, y=4.08, w=1.13, h=0.14, size=7.1, color=MUTED, italic=True, align="center", name="Role_Maps_Label")
    _add_svg_box(parts, left=2.02, top=3.62, width=1.86, height=0.64, fill=ORANGE_LIGHT, line=ORANGE, radius=0.08, dashed=True, name="Training_Only_Anchors")
    _add_svg_text(parts, "training-only anchors", x=2.13, y=3.73, w=1.64, h=0.15, size=8.2, color=ORANGE, bold=True, align="center", name="Training_Only_Anchors_Title")
    _add_svg_text(parts, "A_B=W_L(V),  A_S=W_H(V)", x=2.10, y=3.96, w=1.70, h=0.17, size=8.4, color=ORANGE, align="center", font="Cambria Math", name="Training_Only_Anchors_Formula")
    _add_svg_text(parts, "not available at inference", x=2.13, y=4.18, w=1.64, h=0.13, size=7.1, color=MUTED, italic=True, align="center", name="Training_Only_Anchors_Note")
    _add_svg_line(parts, 2.95, 3.62, 3.58, 3.12, color=ORANGE, width=0.85, dashed=True, arrow=False, name="Anchor_to_RoleHeads")
    _add_svg_line(parts, 3.58, 3.12, 4.24, 2.60, color=ORANGE, width=0.85, dashed=True, arrow=False, name="Anchor_to_CS")

    _add_svg_text(parts, "restricted routing", x=5.87, y=0.86, w=1.25, h=0.16, size=8.1, color=MUTED, italic=True, name="Restricted_Routing_Label")
    _add_svg_box(parts, left=5.92, top=1.25, width=0.76, height=0.48, fill=GREEN_LIGHT, line=GREEN, name="B_Condition_Node")
    _add_svg_text(parts, "C_B", x=5.99, y=1.38, w=0.62, h=0.18, size=12.0, color=GREEN, bold=True, align="center", font="Cambria Math", name="B_Condition_Label")
    _add_svg_box(parts, left=6.88, top=1.12, width=1.42, height=0.74, fill=GREEN_LIGHT, line=GREEN, name="Background_UNet")
    _add_svg_text(parts, "Background U-Net", x=7.00, y=1.26, w=1.18, h=0.17, size=9.0, color=GREEN, bold=True, align="center", name="Background_UNet_Title")
    _add_svg_text(parts, "f_B(C_B)", x=7.13, y=1.53, w=0.92, h=0.17, size=9.2, color=GREEN, align="center", font="Cambria Math", name="Background_UNet_Formula")
    _add_svg_line(parts, 6.68, 1.49, 6.88, 1.49, color=GREEN, width=1.2, name="CB_to_Background")
    _add_svg_line(parts, 8.30, 1.49, 8.52, 1.49, color=GREEN, width=1.2, name="Background_to_Bhat")
    _add_svg_box(parts, left=8.52, top=1.25, width=0.85, height=0.48, fill=GREEN_LIGHT, line=GREEN, name="Predicted_Background")
    _add_svg_text(parts, "B̂", x=8.58, y=1.38, w=0.73, h=0.18, size=11.5, color=GREEN, bold=True, align="center", font="Cambria Math", name="Predicted_Background_Label")
    _add_svg_box(parts, left=8.45, top=1.98, width=1.02, height=0.48, fill=WHITE, line=GREEN, name="Stop_Gradient_Node")
    _add_svg_text(parts, "B̄=sg[B̂]", x=8.50, y=2.11, w=0.92, h=0.18, size=8.0, color=GREEN, bold=True, align="center", font="Cambria Math", name="Stop_Gradient_Label")
    _add_svg_line(parts, 8.95, 1.73, 8.95, 1.98, color=GREEN, width=1.0, name="Bhat_to_StopGradient")
    _add_svg_text(parts, "∇ωB L_R = 0", x=9.55, y=1.99, w=1.08, h=0.18, size=7.7, color=PURPLE, italic=True, font="Cambria Math", name="Stop_Gradient_Constraint")
    _add_svg_box(parts, left=5.92, top=2.52, width=0.76, height=0.48, fill=BLUE_LIGHT, line=BLUE, name="S_Condition_Node")
    _add_svg_text(parts, "C_S", x=5.99, y=2.65, w=0.62, h=0.18, size=12.0, color=BLUE, bold=True, align="center", font="Cambria Math", name="S_Condition_Label")
    _add_svg_line(parts, 5.06, 1.87, 5.92, 1.49, color=GREEN, width=1.0, name="Contract_CB_to_Responsibility")
    _add_svg_line(parts, 5.06, 2.60, 5.92, 2.76, color=BLUE, width=1.0, name="Contract_CS_to_Responsibility")
    _add_svg_box(parts, left=9.67, top=2.46, width=1.35, height=0.62, fill=BLUE_LIGHT, line=BLUE, name="Residual_Context")
    _add_svg_text(parts, "c=[C_S;ψ_B(B̄)]", x=9.76, y=2.61, w=1.17, h=0.18, size=8.1, color=BLUE, bold=True, align="center", font="Cambria Math", name="Residual_Context_Formula")
    _add_svg_line(parts, 8.95, 2.46, 9.67, 2.77, color=GREEN, width=1.0, name="Bbar_to_Context")
    _add_svg_line(parts, 6.68, 2.76, 9.67, 2.76, color=BLUE, width=1.1, name="CS_to_Context")
    _add_svg_box(parts, left=10.13, top=3.24, width=2.05, height=0.75, fill=BLUE_LIGHT, line=BLUE, name="Residual_FM_Block")
    _add_svg_text(parts, "velocity-space FiLM", x=10.26, y=3.36, w=1.80, h=0.16, size=9.1, color=BLUE, bold=True, align="center", name="Residual_FM_Title")
    _add_svg_text(parts, "Residual Flow Matching", x=10.28, y=3.60, w=1.76, h=0.16, size=8.7, color=BLUE, align="center", name="Residual_FM_Subtitle")
    _add_svg_line(parts, 10.35, 3.08, 10.35, 3.24, color=BLUE, width=1.1, name="Context_to_FM")
    transport_specs = [("noise", "ξ", 8.00, 3.39, "Transport_Noise"), ("rt", "R_t", 8.65, 3.39, "Transport_Rt"), ("rhat", "R_{B̂}", 9.30, 3.39, "Transport_Rhat")]
    for key, label, left, top, name in transport_specs:
        _add_svg_image(parts, TILES[key], left=left, top=top, width=0.48, height=0.48, name=f"{name}_Tile")
        _add_svg_text(parts, label, x=left - 0.02, y=top + 0.50, w=0.52, h=0.14, size=7.4, color=BLUE, align="center", font="Cambria Math", name=f"{name}_Label")
    _add_svg_line(parts, 8.49, 3.63, 8.65, 3.63, color=BLUE, width=0.8, name="Noise_to_Rt")
    _add_svg_line(parts, 9.14, 3.63, 9.30, 3.63, color=BLUE, width=0.8, name="Rt_to_Rhat")
    _add_svg_text(parts, "R_t=(1-t)ξ+tR_{B̂}", x=7.94, y=4.09, w=1.90, h=0.17, size=8.0, color=BLUE, align="center", font="Cambria Math", name="Transport_Interpolation_Formula")
    _add_svg_text(parts, "u_t=R_{B̂}-ξ", x=9.72, y=4.09, w=1.10, h=0.17, size=8.0, color=BLUE, align="center", font="Cambria Math", name="Transport_Target_Formula")
    _add_svg_box(parts, left=6.00, top=3.76, width=1.68, height=0.55, fill=ORANGE_LIGHT, line=ORANGE, radius=0.08, dashed=True, name="Residual_Target_TrainingOnly")
    _add_svg_text(parts, "V (training only)", x=6.10, y=3.86, w=1.48, h=0.15, size=8.0, color=ORANGE, bold=True, align="center", name="Residual_Target_V_Label")
    _add_svg_text(parts, "R_{B̂}=V-B̄", x=6.12, y=4.08, w=1.44, h=0.16, size=8.3, color=ORANGE, align="center", font="Cambria Math", name="Residual_Target_Formula")
    _add_svg_line(parts, 7.68, 4.03, 9.30, 4.03, color=ORANGE, width=0.85, dashed=True, arrow=False, name="Target_to_Residual_Transport")
    _add_svg_box(parts, left=11.35, top=1.36, width=0.43, height=0.43, fill=WHITE, line=TEXT, radius=0, name="Composition_Plus")
    _add_svg_text(parts, "+", x=11.35, y=1.42, w=0.43, h=0.20, size=15.0, bold=True, align="center", name="Composition_Plus_Label")
    _add_svg_line(parts, 9.37, 1.49, 11.35, 1.58, color=GREEN, width=1.1, name="Bhat_to_Composition")
    _add_svg_line(parts, 11.13, 3.60, 11.55, 1.79, color=BLUE, width=1.1, name="Rhat_to_Composition")
    _add_svg_box(parts, left=11.82, top=1.24, width=1.18, height=0.68, fill=WHITE, line=TEXT, name="Final_Output")
    _add_svg_text(parts, "V̂ = B̂ + R̂", x=11.90, y=1.46, w=1.02, h=0.22, size=8.2, bold=True, align="center", font="Cambria Math", name="Final_Output_Formula")
    _add_svg_line(parts, 11.78, 1.58, 11.82, 1.58, color=TEXT, width=1.0, name="Composition_to_Output")
    _add_svg_text(parts, "V̂-V=R̂-R_{B̂}", x=11.16, y=2.04, w=1.75, h=0.18, size=8.0, color=PURPLE, align="center", font="Cambria Math", name="Composition_Consistency")
    _add_svg_box(parts, left=10.72, top=4.14, width=2.35, height=0.56, fill=PURPLE_LIGHT, line=PURPLE, name="Transport_Energy_Condition")
    _add_svg_text(parts, "E||u_R||² < E||u_V||²", x=10.84, y=4.20, w=2.10, h=0.16, size=10.4, color=PURPLE, align="center", font="Cambria Math", name="Transport_Energy_Formula_1")
    _add_svg_text(parts, "iff E||V-B̂||² < E||V||²", x=10.84, y=4.39, w=2.10, h=0.16, size=10.4, color=PURPLE, align="center", font="Cambria Math", name="Transport_Energy_Formula_2")
    parts.append("</svg>")
    return "\n".join(parts)


def _make_pptx(output_path: Path) -> dict[str, int | bool]:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(WHITE)
    _draw_ppt(slide)
    presentation.save(output_path)

    with zipfile.ZipFile(output_path) as archive:
        media_count = sum(name.startswith("ppt/media/") for name in archive.namelist())
    shape_count = len(slide.shapes)
    return {
        "slide_count": len(presentation.slides),
        "shape_count": shape_count,
        "picture_count": media_count,
        "all_text_arrows_native_editable": True,
    }


def _export_svg_assets(output_dir: Path) -> dict[str, str]:
    svg_path = output_dir / "method_overview.svg"
    pdf_path = output_dir / "method_overview.pdf"
    png_path = output_dir / "method_overview_preview.png"
    svg_path.write_text(_draw_svg(), encoding="utf-8")
    subprocess.run(["rsvg-convert", "-f", "pdf", "-o", str(pdf_path), str(svg_path)], check=True)
    subprocess.run(["rsvg-convert", "-w", "3200", "-h", "1200", "-o", str(png_path), str(svg_path)], check=True)
    with Image.open(png_path) as image:
        rgb = image.convert("RGB")
        rgb.save(png_path, dpi=(300, 300))
    return {"svg": str(svg_path), "pdf": str(pdf_path), "png": str(png_path)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()
    _check_assets()
    output_dir = args.output_dir if args.output_dir.is_absolute() else REPO_ROOT / args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)

    pptx_path = output_dir / "method_overview_editable.pptx"
    ppt_info = _make_pptx(pptx_path)
    vector_info = _export_svg_assets(output_dir)
    manifest = {
        "figure": "PD-BG-RFM method overview",
        "source_pptx": str(SOURCE_PPTX),
        "canvas_inches": [SLIDE_W, SLIDE_H],
        "outputs": {"pptx": str(pptx_path), **vector_info},
        "pptx_validation": ppt_info,
        "asset_map": {key: str(path) for key, path in TILES.items()},
        "semantic_colors": {
            "background": f"#{GREEN}",
            "structure": f"#{BLUE}",
            "training_only": f"#{ORANGE}",
            "theory": f"#{PURPLE}",
        },
        "formula_contract": [
            "p_PD(b,r|XΩ) = δ[b=f_B(C_B)] pθ(r|C_S,b)",
            "B̄=sg[B̂]",
            "R_{B̂}=V-B̄",
            "R_t=(1-t)ξ+tR_{B̂}",
            "V̂=B̂+R̂",
            "V̂-V=R̂-R_{B̂}",
        ],
        "training_only_is_dashed_orange": True,
        "forbidden_legacy_terms_absent": True,
    }
    manifest_path = output_dir / "method_overview_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
